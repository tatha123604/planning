from __future__ import annotations

from copy import copy
from datetime import date, datetime
from io import BytesIO
from os import PathLike
import re

from openpyxl import load_workbook

from processor import coerce_report_date, format_report_date

TARGET_SHEET_TITLE = "NON CONT DUTY SIGNON SIGNOFF"
SIGN_ON_TITLE = "NON CONTINUOUS DUTY SIGN_ON"
SIGN_OFF_TITLE = "NON CONTINUOUS DUTY SIGN_OFF"


def _read_bytes(file_obj) -> bytes:
    if isinstance(file_obj, (bytes, bytearray)):
        return bytes(file_obj)
    if isinstance(file_obj, (str, PathLike)):
        with open(file_obj, "rb") as handle:
            return handle.read()
    if hasattr(file_obj, "getvalue"):
        return file_obj.getvalue()
    if hasattr(file_obj, "read"):
        pos = None
        if hasattr(file_obj, "tell"):
            pos = file_obj.tell()
        data = file_obj.read()
        if hasattr(file_obj, "seek") and pos is not None:
            file_obj.seek(pos)
        return data
    raise TypeError("Unsupported file input")


def _as_stream(file_obj) -> BytesIO:
    return BytesIO(_read_bytes(file_obj))


def _normalize_header(value) -> str:
    text = str(value or "").replace("\n", " ").replace(".", " ")
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


def _normalize_label(value) -> str:
    return _normalize_header(value).replace(" ", "")


def _display_value(value) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d-%m-%y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header_row(ws) -> int:
    for row in range(1, min(ws.max_row, 12) + 1):
        values = [_normalize_header(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        if "CREW ID" in values and "CREW NAME" in values:
            return row
    raise ValueError("Could not find crew header row in uploaded workbook.")


def _detect_section(ws, header_row: int) -> str:
    heading_parts = []
    for row in range(1, min(header_row, 4) + 1):
        heading_parts.extend(
            str(ws.cell(row, col).value or "") for col in range(1, min(ws.max_column, 4) + 1)
        )
    heading_text = " ".join(heading_parts).upper()
    if "SIGN_OFF" in heading_text or "SIGNOFF" in heading_text:
        return "sign_off"

    header_text = " ".join(
        _normalize_header(ws.cell(header_row, col).value) for col in range(1, ws.max_column + 1)
    )
    if "SIGNOFF" in header_text or "SIGN OFF" in header_text:
        return "sign_off"
    return "sign_on"


def _collect_rows(ws, header_row: int, section: str) -> list[dict]:
    header_map = {}
    for col in range(1, ws.max_column + 1):
        key = _normalize_header(ws.cell(header_row, col).value)
        if key:
            header_map[key] = col

    required = ["CREW ID", "CREW NAME", "DESIG", "SUP ID", "DUTY TYPE"]
    missing = [name for name in required if name not in header_map]
    if missing:
        raise ValueError(f"Missing source columns: {', '.join(missing)}")

    station_keys = ["SIGNON STTN", "SIGN ON STTN"] if section == "sign_on" else ["SIGNOFF STTN", "SIGN OFF STTN"]
    time_keys = ["SIGNON TIME", "SIGN ON TIME"] if section == "sign_on" else ["SIGNOFF TIME", "SIGN OFF TIME"]
    route_keys = ["TO STN"] if section == "sign_on" else ["FROM STN", "TO STN"]

    def pick(*candidates):
        for name in candidates:
            if name in header_map:
                return header_map[name]
        return None

    station_col = pick(*station_keys)
    time_col = pick(*time_keys)
    entry_col = pick("ENTRY POINT", "ENTRYPOINT")
    train_col = pick("TRAIN NO")
    loco_col = pick("LOCO NO")
    route_col = pick(*route_keys)
    reason_col = pick("REASON")

    rows = []
    for row_no in range(header_row + 1, ws.max_row + 1):
        crew_id = _display_value(ws.cell(row_no, header_map["CREW ID"]).value)
        crew_name = _display_value(ws.cell(row_no, header_map["CREW NAME"]).value)
        if not crew_id and not crew_name:
            continue
        rows.append(
            {
                "SNO.": len(rows) + 1,
                "CREW ID": crew_id,
                "CREW NAME": crew_name,
                "DESIG.": _display_value(ws.cell(row_no, header_map["DESIG"]).value),
                "STATION": _display_value(ws.cell(row_no, station_col).value) if station_col else "",
                "EVENT TIME": _display_value(ws.cell(row_no, time_col).value) if time_col else "",
                "SUP ID": _display_value(ws.cell(row_no, header_map["SUP ID"]).value),
                "ENTRY POINT": _display_value(ws.cell(row_no, entry_col).value) if entry_col else "",
                "TRAIN NO.": _display_value(ws.cell(row_no, train_col).value) if train_col else "",
                "LOCO NO.": _display_value(ws.cell(row_no, loco_col).value) if loco_col else "",
                "DUTY TYPE": _display_value(ws.cell(row_no, header_map["DUTY TYPE"]).value),
                "ROUTE STN": _display_value(ws.cell(row_no, route_col).value) if route_col else "",
                "REASON": _display_value(ws.cell(row_no, reason_col).value) if reason_col else "",
            }
        )
    return rows


def parse_non_continuous_source(source_file) -> tuple[str, list[dict]]:
    workbook = load_workbook(_as_stream(source_file), data_only=True)
    ws = workbook[workbook.sheetnames[0]]
    header_row = _find_header_row(ws)
    section = _detect_section(ws, header_row)
    return section, _collect_rows(ws, header_row, section)


def _copy_cell_style(source_cell, target_cell) -> None:
    target_cell._style = copy(source_cell._style)
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


def _apply_row_style(ws, source_row: int, target_row: int, end_col: int) -> None:
    for col in range(1, end_col + 1):
        _copy_cell_style(ws.cell(source_row, col), ws.cell(target_row, col))
    if ws.row_dimensions[source_row].height is not None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _find_row_by_text(ws, text: str, start_row: int = 1) -> int:
    target = _normalize_label(text)
    for row in range(start_row, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = _normalize_label(ws.cell(row, col).value)
            if value and target in value:
                return row
    raise ValueError(f"Could not find row for '{text}' in template.")


def _find_template_sheet(workbook, report_date_value=None):
    report_date = coerce_report_date(report_date_value)
    month_tokens = []
    if report_date:
        month_tokens = [
            report_date.strftime("%B").upper(),
            report_date.strftime("%b").upper(),
        ]

    best_sheet = None
    best_score = -1
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        try:
            _find_row_by_text(ws, SIGN_ON_TITLE)
            has_sign_on = True
        except ValueError:
            has_sign_on = False
        try:
            _find_row_by_text(ws, SIGN_OFF_TITLE)
            has_sign_off = True
        except ValueError:
            has_sign_off = False

        score = 0
        if has_sign_on:
            score += 2
        if has_sign_off:
            score += 2
        upper_name = sheet_name.upper()
        if any(token and token in upper_name for token in month_tokens):
            score += 3

        if score > best_score:
            best_sheet = ws
            best_score = score

    if best_sheet is None:
        raise ValueError("Template workbook must contain at least one sheet.")
    return best_sheet


def _clear_range(ws, start_row: int, end_row: int, end_col: int) -> None:
    if end_row < start_row:
        return
    for row in range(start_row, end_row + 1):
        for col in range(1, end_col + 1):
            ws.cell(row=row, column=col).value = None


def _write_section(
    ws,
    rows: list[dict],
    title_row: int,
    title_text: str,
    station_heading: str,
    route_heading: str,
    block_end: int | None = None,
) -> None:
    header_row = title_row + 1
    data_start = title_row + 2
    end_col = 13
    if block_end is None:
        block_end = max(ws.max_row, data_start + len(rows) + 5)

    ws.cell(title_row, 1).value = title_text
    ws.cell(header_row, 5).value = station_heading
    ws.cell(header_row, 6).value = "SIGNON TIME" if "SIGN_ON" in title_text else "SIGNOFF TIME"
    ws.cell(header_row, 12).value = route_heading

    _clear_range(ws, data_start, block_end, end_col)

    template_row = data_start if data_start <= ws.max_row else header_row
    for offset, record in enumerate(rows):
        target_row = data_start + offset
        if target_row > ws.max_row:
            ws.insert_rows(target_row)
        _apply_row_style(ws, template_row, target_row, end_col)
        values = [
            record["SNO."],
            record["CREW ID"],
            record["CREW NAME"],
            record["DESIG."],
            record["STATION"],
            record["EVENT TIME"],
            record["SUP ID"],
            record["ENTRY POINT"],
            record["TRAIN NO."],
            record["LOCO NO."],
            record["DUTY TYPE"],
            record["ROUTE STN"],
            record["REASON"],
        ]
        for col_index, value in enumerate(values, start=1):
            ws.cell(target_row, col_index).value = value


def build_non_continuous_workbook(
    sign_on_rows: list[dict],
    sign_off_rows: list[dict],
    template_file,
    report_date_value=None,
) -> BytesIO:
    workbook = load_workbook(_as_stream(template_file))
    if not workbook.sheetnames:
        raise ValueError("Template workbook must contain at least one sheet.")

    if TARGET_SHEET_TITLE in workbook.sheetnames:
        del workbook[TARGET_SHEET_TITLE]
    ws = workbook.copy_worksheet(_find_template_sheet(workbook, report_date_value))
    ws.title = TARGET_SHEET_TITLE

    sign_on_title_row = _find_row_by_text(ws, SIGN_ON_TITLE, 1)
    sign_off_title_row = _find_row_by_text(ws, SIGN_OFF_TITLE, sign_on_title_row + 1)
    sign_on_data_start = sign_on_title_row + 2
    sign_on_capacity = max(0, sign_off_title_row - sign_on_data_start)
    extra_sign_on_rows = max(0, len(sign_on_rows) - sign_on_capacity)
    if extra_sign_on_rows:
        ws.insert_rows(sign_off_title_row, extra_sign_on_rows)
        sign_off_title_row += extra_sign_on_rows

    _write_section(
        ws,
        sign_on_rows,
        sign_on_title_row,
        SIGN_ON_TITLE,
        "SIGNON STTN",
        "To STN",
        block_end=sign_off_title_row - 1,
    )
    _write_section(
        ws,
        sign_off_rows,
        sign_off_title_row,
        SIGN_OFF_TITLE,
        "SIGNOFF STTN",
        "From STN",
    )

    report_date = format_report_date(report_date_value)
    if report_date:
        ws.cell(1, 15).value = f"Report Date: {report_date}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
