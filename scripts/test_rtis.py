"""RTIS import and route checks against an isolated database.

Run: python scripts/test_rtis.py [path/to/RTIS_export.xlsx]
"""
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook, Workbook
from pypdf import PdfReader
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, create_engine, select

from src.rtis import (HEADERS, RtisEvent, RtisUpload, get_session,
                      import_rtis, parse_rtis, router)
from src.rtis_train_models import RtisTrainModel, parse_train_model, save_train_model
from src.rtis_homes import RtisHomeModel, _polygon_centroid, event_home_details, parse_home_model, save_home_model


def model_workbook(rows=None):
    book = Workbook()
    sheet = book.active
    sheet.append(['Train No', 'NAME OF TRAIN', 'Train No', 'HQ OF CREW'])
    for row in rows or [(12377, 'PADATIK', 12378, 'SDAH'),
                        ('00441', 'Special\nTrain', '00442', ' KOAA '),
                        (13185, 'Excluded', 13186, 'ASN')]:
        sheet.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def home_workbook():
    book = Workbook()
    sheet = book.active
    sheet.append(['Station', 'DIRN', 'Station DIRN', 'TYPE', 'Latitude', 'Longitude'])
    sheet.append(['BP', 'UP MAIN LINE', 'BP-UP MAIN LINE', 'Home', 22.1, 88.1])
    sheet.append(['BP', 'DN MAIN LINE', 'BP-DN MAIN LINE', 'I/Home', 22.2, 88.2])
    sheet.append(['BP', 'UP LOOP LINE', 'BP-UP LOOP LINE', 'Home', 22.11, 88.11])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()

REFERENCE = Path(sys.argv.pop()) if len(sys.argv) > 1 and sys.argv[-1].endswith('.xlsx') else None


def workbook(division='SDAH', speed='0', event='H', time='2026-09-14 10:00:00', serial='1', formula=False, train='00441'):
    from xml.sax.saxutils import escape
    values = [serial, '12', '22364', '22.76', '88.37', 'BP', time, event,
              speed, division, '2026-09-14 10:00:02', train, '2026-09-14']
    def row(number, items):
        cells = []
        for i, value in enumerate(items):
            inner = '<f>1+1</f><v>2</v>' if formula and number == 2 and i == 8 else f'<is><t>{escape(value)}</t></is>'
            cells.append(f'<c r="{chr(65+i)}{number}" t="inlineStr">{inner}</c>')
        return f'<row r="{number}">' + ''.join(cells) + '</row>'
    stream = BytesIO()
    with ZipFile(stream, 'w') as archive:
        archive.writestr('xl/workbook.xml', '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Events" sheetId="1" r:id="rId1"/></sheets></workbook>')
        archive.writestr('xl/_rels/workbook.xml.rels', '<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        archive.writestr('xl/worksheets/sheet1.xml', '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + row(1, HEADERS) + row(2, values) + '</sheetData></worksheet>')
    return stream.getvalue()


class RtisTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        RtisUpload.__table__.create(self.engine)
        RtisEvent.__table__.create(self.engine)
        RtisTrainModel.__table__.create(self.engine)
        RtisHomeModel.__table__.create(self.engine)
        self.session = Session(self.engine)
        self.app = FastAPI()
        self.app.include_router(router)
        self.app.dependency_overrides[get_session] = lambda: self.session
        self.client = TestClient(self.app)

    def tearDown(self):
        self.client.close()
        self.session.close()
        self.engine.dispose()

    def test_train_model_pairs_hq_names_and_validation(self):
        mapping, count = parse_train_model(model_workbook())
        self.assertEqual(count, 2)
        self.assertEqual(set(mapping), {'12377', '12378', '00441', '00442'})
        self.assertEqual(mapping['12378']['name'], 'PADATIK')
        self.assertEqual(mapping['00442'], {'name': 'Special Train', 'hq': 'KOAA'})
        for rows in [[(1, 'Excluded', 2, 'ASN')], [(1, '', 2, 'SDAH')],
                     [('ABC123', 'Bad', 2, 'SDAH')],
                     [(1, 'One', 2, 'SDAH'), (1, 'Other', 3, 'KOAA')]]:
            with self.assertRaises(ValueError):
                parse_train_model(model_workbook(rows))
        with self.assertRaises(ValueError):
            parse_train_model(b'bad file')

    def test_home_model_maps_directions_and_distance(self):
        self.assertEqual(_polygon_centroid([(0, 0), (4, 0), (0, 4)]), (4 / 3, 4 / 3))
        mapping, signals = parse_home_model(home_workbook())
        self.assertEqual(signals, 3)
        self.assertEqual(mapping['BP|J']['homes'][0]['type'], 'Home')
        self.assertEqual(mapping['BP|K']['homes'][0]['type'], 'I/Home')
        from src.rtis_homes import enrich_home_model
        mapped = enrich_home_model(mapping, {'BP': [(88.0, 22.0), (88.0, 22.01), (88.01, 22.01)]})
        self.assertEqual(mapped, 2)
        self.assertEqual(event_home_details('bp', 'J', mapping)[0], 'UP HOME')
        self.assertGreater(event_home_details('BP', 'J', mapping)[1], 0)
        self.assertEqual(event_home_details('BP', 'H', mapping), ('', ''))
        save_home_model(self.session, 'fsd.xlsx', home_workbook())
        refreshed = save_home_model(self.session, 'same-fsd.xlsx', home_workbook(), {'BP': [(88.0, 22.0), (88.0, 22.01), (88.01, 22.01)]})
        self.assertEqual(refreshed.mapped_station_count, 2)
        import_rtis(self.session, 'bp.xlsx', workbook(event='J', time='2026-09-14 10:00:00'), 'SDAH')
        page = self.client.get('/rtis?day=2026-09-14')
        self.assertIn('UP HOME', page.text)
        self.assertIn('Home distance (m)', page.text)
        self.assertIn('FSD Home signal coordinates', page.text)
        self.assertIn('Station DIRN', page.text)
        self.assertIn('UP MAIN LINE', page.text)
        filtered = self.client.get('/rtis?home_station=bp')
        self.assertEqual(filtered.context['home_station'], 'BP')
        self.assertIn('FSD Station name / code', filtered.text)
        home_book = load_workbook(BytesIO(self.client.get('/rtis/home-model.xlsx?home_station=BP').content))
        self.assertEqual(home_book.active.max_row, 4)
        self.assertEqual(home_book.active['A2'].value, 'BP')
        self.assertEqual(home_book.active['I1'].value, 'Linear distance (m)')
        home_book.close()
        self.assertEqual(self.client.get('/rtis/home-model.xlsx?home_station=NOTFOUND').status_code, 200)
        book = load_workbook(BytesIO(self.client.get('/rtis/analysis.xlsx?day=2026-09-14').content))
        self.assertEqual([cell.value for cell in book.active[1]][-2:], ['Home Signal', 'Home Distance (m)'])
        book.close()

    def test_train_model_upload_filter_export_and_reuse(self):
        for index, (train, division) in enumerate([('12377', 'SDAH'), ('12378', 'HWH'),
                                ('00441', 'ASN'), ('13185', 'SDAH'),
                                ('912377', 'SDAH'), ('GO123', 'SDAH')]):
            import_rtis(self.session, train + '.xlsx', workbook(train=train, division=division,
                        time=f'2026-09-14 10:00:0{index}'), division)
        content = model_workbook()
        response = self.client.post('/rtis/train-model/upload', data={'division': 'ALL'},
                                    files={'model_file': ('model.xlsx', content)})
        self.assertEqual(response.status_code, 200)
        self.assertIn('3 matching events', response.text)
        self.assertIn('Special Train', response.text)
        self.assertIn('Train Name', response.text)
        model = self.session.exec(select(RtisTrainModel)).one()
        self.assertEqual(save_train_model(self.session, 'copy.xlsx', content).id, model.id)
        query = f'division=ALL&model_id={model.id}'
        self.assertIn('1 matching events', self.client.get('/rtis?' + query + '&division=HWH').text)
        self.assertIn('2 matching events', self.client.get('/rtis?' + query + '&train_no=1237').text)
        self.assertIn('0 matching events', self.client.get('/rtis?' + query + '&speed=30').text)
        named = self.client.get('/rtis?' + query + '&train_name=PADATIK')
        self.assertEqual(named.context['total'], 2)
        self.assertIn('Search train name', named.text)
        self.assertIn('PADATIK', named.text)
        self.assertEqual(self.client.get('/rtis?' + query + '&train_name=NO-SUCH-NAME').context['total'], 0)
        self.assertIn('5 matching events', self.client.get('/rtis?division=ALL').text)
        self.assertIn('1 matching events', self.client.get('/rtis?division=ALL&analysis=goods').text)
        result = self.client.get('/rtis/analysis.xlsx?' + query)
        self.assertEqual(result.status_code, 200)
        book = load_workbook(BytesIO(result.content))
        rows = list(book.active.values)
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0][-2:], ('Train Name', 'HQ OF CREW'))
        names = {row[4]: row[-2] for row in rows[1:]}
        self.assertEqual(names, {'12377': 'PADATIK', '12378': 'PADATIK', '00441': 'Special Train'})
        book.close()
        self.assertEqual(self.client.get('/rtis?model_id=999').status_code, 404)
        self.assertEqual(self.client.get('/rtis?' + query + '&analysis=goods').status_code, 400)
        invalid = self.client.post('/rtis/train-model/upload', files={'model_file': ('bad.xlsx', b'bad')})
        self.assertIn('Model not saved', invalid.text)
        self.assertEqual(len(self.session.exec(select(RtisTrainModel)).all()), 1)

    def test_division_validation_and_atomic_failure(self):
        for division in ('SDAH', 'HWH', 'ASN', 'MLDT'):
            import_rtis(self.session, 'report.xlsx', workbook(division=division), division)
        with self.assertRaisesRegex(ValueError, 'expected HWH'):
            import_rtis(self.session, 'bad.xlsx', workbook(), 'HWH')
        self.assertEqual(len(self.session.exec(select(RtisUpload)).all()), 4)

    def test_duplicate_file_and_overlapping_exports(self):
        content = workbook()
        import_rtis(self.session, 'report.xlsx', content, 'SDAH')
        self.assertIn('Already uploaded', import_rtis(self.session, 'copy.xlsx', content, 'SDAH'))
        self.assertIn('0 new events', import_rtis(self.session, 'overlap.xlsx', workbook(serial='88'), 'SDAH'))
        self.assertEqual(len(self.session.exec(select(RtisEvent)).all()), 1)

    def test_rtis_retains_only_latest_two_event_dates(self):
        for day in ('2026-09-13', '2026-09-14', '2026-09-15'):
            import_rtis(self.session, f'{day}.xlsx', workbook(time=f'{day} 10:00:00'), 'SDAH')
        events = self.session.exec(select(RtisEvent).order_by(RtisEvent.event_time)).all()
        self.assertEqual([event.event_time.date().isoformat() for event in events], ['2026-09-14', '2026-09-15'])
        uploads = self.session.exec(select(RtisUpload).order_by(RtisUpload.filename)).all()
        self.assertEqual([upload.filename for upload in uploads], ['2026-09-14.xlsx', '2026-09-15.xlsx'])
        self.assertIn('1 matching events', self.client.get('/rtis?day=2026-09-14').text)
        self.assertIn('1 matching events', self.client.get('/rtis?day=2026-09-15').text)
        self.assertIn('0 matching events', self.client.get('/rtis?day=2026-09-13').text)

    def test_identifiers_blank_zero_and_excel_dates(self):
        row = parse_rtis(workbook(), 'SDAH')[0]
        self.assertEqual(row['train'], '00441')
        self.assertEqual(row['speed'], 0)
        self.assertIsNone(parse_rtis(workbook(speed=''), 'SDAH')[0]['speed'])
        self.assertEqual(parse_rtis(workbook(time='46279.5'), 'SDAH')[0]['event_time'].isoformat(), '2026-09-14T12:00:00')

    def test_invalid_files_are_not_saved(self):
        for content in (b'not a workbook', workbook(speed='nan'), workbook(time='bad'), workbook(formula=True)):
            with self.assertRaises(ValueError):
                import_rtis(self.session, 'bad.xlsx', content, 'SDAH')
        self.assertEqual(self.session.exec(select(RtisUpload)).all(), [])

    def test_upload_filters_history_download_and_errors(self):
        files = [('files', ('one.xlsx', workbook(event='H'))), ('files', ('two.xlsx', workbook(event='J')))]
        response = self.client.post('/rtis/upload', data={'division': 'SDAH'}, files=files)
        self.assertEqual(response.status_code, 200)
        self.assertIn('2 matching events', response.text)
        self.assertIn('00441', response.text)
        self.assertIn('one.xlsx', response.text)
        self.assertIn('1 matching events', self.client.get('/rtis?event=J&day=2026-09-14').text)
        self.assertIn('0 matching events', self.client.get('/rtis?day=2026-09-15').text)
        self.assertIn('0 matching events', self.client.get('/rtis?division=HWH').text)
        self.assertEqual(self.client.get('/rtis?division=BAD').status_code, 400)
        self.assertEqual(self.client.get('/rtis?day=invalid').status_code, 400)
        original = self.session.get(RtisUpload, 1).content
        self.assertEqual(self.client.get('/rtis/uploads/1/download').content, original)
        self.assertEqual(self.client.get('/rtis/uploads/999/download').status_code, 404)
        failure = self.client.post('/rtis/upload', data={'division':'HWH'}, files={'files':('bad.xlsx', workbook())})
        self.assertIn('Not saved:', failure.text)

    def test_undo_upload_removes_only_that_file_and_events(self):
        import_rtis(self.session, 'one.xlsx', workbook(serial='1'), 'SDAH')
        import_rtis(self.session, 'two.xlsx', workbook(serial='2', time='2026-09-14 11:00:00'), 'SDAH')
        response = self.client.post('/rtis/uploads/1/delete', data={'division': 'SDAH', 'analysis': 'passenger'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('Undone one.xlsx: 1 events removed.', response.text)
        self.assertEqual(len(self.session.exec(select(RtisUpload)).all()), 1)
        self.assertEqual(len(self.session.exec(select(RtisEvent)).all()), 1)
        self.assertEqual(self.client.get('/rtis/uploads/1/download').status_code, 404)
        self.assertEqual(self.client.post('/rtis/uploads/999/delete').status_code, 404)

    def test_pagination_and_nonfocus_retention(self):
        for i in range(101):
            import_rtis(self.session, f'{i}.xlsx', workbook(time=f'2026-09-14 {i//60:02}:{i%60:02}:00'), 'SDAH')
        import_rtis(self.session, 'arrival.xlsx', workbook(event='A'), 'SDAH')
        self.assertEqual(len(self.session.exec(select(RtisEvent)).all()), 102)
        self.assertIn('101 matching events', self.client.get('/rtis').text)
        self.assertIn('Page 2 of 2', self.client.get('/rtis?page=2').text)

    def test_speed_sort_is_applied_before_pagination_and_export(self):
        for index in range(101):
            import_rtis(self.session, f'sort-{index}.xlsx', workbook(
                serial=str(index + 1), speed=str(index % 70), time=f'2026-09-14 10:{index // 60:02}:{index % 60:02}'), 'SDAH')
        first = self.client.get('/rtis?sort_by=speed&sort_order=asc&page=1')
        second = self.client.get('/rtis?sort_by=speed&sort_order=asc&page=2')
        self.assertEqual(first.context['rows'][0].speed, 0)
        self.assertEqual(second.context['rows'][0].speed, 69)
        self.assertEqual(first.context['sort_by'], 'speed')
        self.assertEqual(second.context['sort_order'], 'asc')
        book = load_workbook(BytesIO(self.client.get('/rtis/analysis.xlsx?sort_by=speed&sort_order=desc').content))
        speeds = [row[6].value for row in book.active.iter_rows(min_row=2)]
        self.assertEqual(speeds[0], 69)
        self.assertEqual(speeds[-1], 0)
        book.close()

    @unittest.skipUnless(REFERENCE, 'No reference workbook supplied')
    def test_actual_export_with_invalid_fills(self):
        content = REFERENCE.read_bytes()
        records = parse_rtis(content, 'SDAH')
        self.assertEqual(Counter(r['event_type'] for r in records), {'A':377, 'D':419, 'H':1464, 'J':940, 'K':933, 'R':923})
        import_rtis(self.session, REFERENCE.name, content, 'SDAH')
        response = self.client.get('/rtis?day=2026-09-14')
        self.assertEqual(response.status_code, 200)
        passenger_count = sum(r['event_type'] in ('H','J','K') and r['train'].isascii() and r['train'].isdigit() for r in records)
        self.assertIn(f'{passenger_count} matching events', response.text)

    def test_train_type_rule_and_selection_persistence(self):
        trains = ['12345', '00441', 'AB123', '123ab', '', 'GOODS', '12-34', 'A/12']
        for i, train in enumerate(trains):
            import_rtis(self.session, f'{i}.xlsx', workbook(train=train, time=f'2026-09-14 10:{i:02}:00'), 'SDAH')
        passenger = self.client.get('/rtis')
        self.assertIn('2 matching events', passenger.text)
        goods = self.client.get('/rtis?analysis=goods&day=2026-09-14')
        self.assertIn('3 matching events', goods.text)
        self.assertIn('Goods Train Analysis — SDAH', goods.text)
        self.assertIn('name="analysis" value="goods"', goods.text)
        self.assertIn('division=HWH&day=2026-09-14&amp;event=HJK&amp;analysis=goods', goods.text)
        self.assertIn('3 matching events', self.client.get('/rtis?view=unclassified').text)
        self.assertEqual(self.client.get('/rtis?analysis=invalid').status_code, 400)
        uploaded = self.client.post('/rtis/upload', data={'division':'SDAH', 'analysis':'goods'},
                                    files={'files': ('goods.xlsx', workbook(train='BC123', event='K'))})
        self.assertIn('analysis=goods', str(uploaded.url))
        self.assertIn('4 matching events', uploaded.text)

    def test_speed_thresholds_date_and_train_type(self):
        for i, speed in enumerate(['', '0', '29.9', '30', '39.9', '40', '49.9', '50', '60']):
            import_rtis(self.session, f'speed{i}.xlsx', workbook(speed=speed, time=f'2026-09-14 10:{i:02}:00'), 'SDAH')
        import_rtis(self.session, 'nextday.xlsx', workbook(speed='60', time='2026-09-15 00:00:00'), 'SDAH')
        import_rtis(self.session, 'goods.xlsx', workbook(speed='50', train='AB123', event='K'), 'SDAH')
        for threshold, count in [('30', 6), ('40', 4), ('50', 2)]:
            response = self.client.get(f'/rtis?speed={threshold}&day=2026-09-14')
            self.assertEqual(response.status_code, 200)
            self.assertIn(f'{count} matching events', response.text)
            self.assertIn(f'value="{threshold}" selected', response.text)
            self.assertIn(f'speed={threshold}', response.context['switch_query'])
            self.assertIn('&analysis=goods', response.text)
        self.assertIn('1 matching events', self.client.get('/rtis?speed=50&analysis=goods&day=2026-09-14').text)
        self.assertIn('9 matching events', self.client.get('/rtis?day=2026-09-14').text)
        self.assertEqual(self.client.get('/rtis?speed=35').status_code, 400)

    def _import_output_examples(self):
        for division, event, train in [('SDAH', 'A', '00441'), ('HWH', 'D', 'AB123'),
                                       ('ASN', 'R', ''), ('MLDT', 'K', '=1+1')]:
            import_rtis(self.session, division + '.xlsx',
                        workbook(division=division, event=event, train=train), division)
        for index in range(101):
            import_rtis(self.session, f'event{index}.xlsx',
                        workbook(serial=str(index + 2), time=f'2026-09-15 {index//60:02}:{index%60:02}:00'), 'SDAH')

    def test_output_combines_all_divisions_types_and_dates_with_pagination(self):
        self._import_output_examples()
        first = self.client.get('/rtis/output')
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.context['headers'], HEADERS)
        self.assertEqual(first.context['total'], 105)
        self.assertEqual(len(first.context['rows']), 100)
        self.assertEqual([row[9] for row in first.context['rows'][:4]], ['SDAH', 'HWH', 'ASN', 'MLDT'])
        self.assertEqual([row[7] for row in first.context['rows'][:4]], ['A', 'D', 'R', 'K'])
        self.assertEqual([row[11] for row in first.context['rows'][:4]], ['00441', 'AB123', '', '=1+1'])
        last = self.client.get('/rtis/output?page=999')
        self.assertEqual(last.context['page'], 2)
        self.assertEqual(len(last.context['rows']), 5)
        self.assertEqual(last.context['last_row'], 105)

    def test_output_excel_contains_all_pages_and_preserves_source_values(self):
        self._import_output_examples()
        response = self.client.get('/rtis/output.xlsx')
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment;', response.headers['content-disposition'])
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(tuple(cell.value for cell in sheet[1]), HEADERS)
        self.assertEqual(sheet.max_row, 106)
        self.assertEqual(sheet['L2'].value, '00441')
        self.assertEqual(sheet['L2'].data_type, 's')
        self.assertEqual(sheet['L5'].value, '=1+1')
        self.assertEqual(sheet['L5'].data_type, 's')
        self.assertEqual(sheet['I2'].value, 0)
        self.assertEqual(sheet['G2'].value, datetime(2026, 9, 14, 10))
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:M106')
        self.assertEqual({sheet.cell(row, 10).value for row in range(2, 107)}, {'SDAH', 'HWH', 'ASN', 'MLDT'})
        book.close()

    def test_output_pdf_contains_all_pages_and_repeats_headers(self):
        self._import_output_examples()
        response = self.client.get('/rtis/output.pdf')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'application/pdf')
        pdf = PdfReader(BytesIO(response.content))
        self.assertGreater(len(pdf.pages), 1)
        for page in pdf.pages:
            text = ' '.join(page.extract_text().split())
            self.assertIn('Rows: 105', text)
            self.assertIn('Page:', text)
            for header in HEADERS:
                self.assertIn(header, text)
        first = pdf.pages[0].extract_text()
        self.assertIn('00441', first)
        self.assertIn('AB123', first)
        self.assertIn('=1+1', first)
        self.assertIn('2026-09-15', pdf.pages[-1].extract_text())
        self.assertGreater(float(pdf.pages[0].mediabox.width), float(pdf.pages[0].mediabox.height))

    def test_filtered_analysis_excel_and_jk_for_both_train_types(self):
        for analysis, train in [('passenger', '00441'), ('goods', 'AB123')]:
            for i in range(101):
                import_rtis(self.session, f'{analysis}{i}.xlsx', workbook(
                    train=train, event='J' if i % 2 else 'K', speed='40',
                    time=f'2026-09-14 {i//60:02}:{i%60:02}:00',
                    division='SDAH' if analysis == 'passenger' else 'HWH'),
                    'SDAH' if analysis == 'passenger' else 'HWH')
            division = 'SDAH' if analysis == 'passenger' else 'HWH'
            for name, kwargs in [('home', {'event':'H'}), ('slow', {'event':'J','speed':'39'}),
                                 ('nextday', {'event':'K','time':'2026-09-15 00:00:00'})]:
                values = dict(train=train, division=division, speed='50')
                values.update(kwargs)
                import_rtis(self.session, name + '.xlsx', workbook(**values), division)
            query = f'division={division}&analysis={analysis}&day=2026-09-14&speed=40&event=JK'
            page = self.client.get('/rtis?' + query)
            self.assertEqual(page.context['total'], 101)
            self.assertEqual({r.event_type for r in page.context['rows']}, {'J','K'})
            self.assertIn('<th>Division Code</th>', page.text)
            self.assertIn('Event date_time<input', page.text)
            self.assertIn('value="JK" selected>J+K', page.text)
            self.assertIn('/rtis/analysis.xlsx?', page.text)
            response = self.client.get('/rtis/analysis.xlsx?' + query)
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(response.content))
            sheet = book.active
            self.assertEqual(sheet.max_row, 102)
            self.assertEqual([c.value for c in sheet[1]], ['Event date_time','Division Code','Station','Event Type','Train Number','Loco No.','Speed'])
            self.assertEqual(sheet['A2'].value, page.context['rows'][0].event_time)
            self.assertEqual(sheet['B2'].value, division)
            self.assertEqual(sheet['E2'].value, train)
            self.assertEqual(sheet['E2'].data_type, 's')
            self.assertEqual(sheet['G2'].value, 40)
            self.assertEqual(sheet.auto_filter.ref, 'A1:G102')
            book.close()
        self.assertEqual(self.client.get('/rtis/analysis.xlsx?event=BAD').status_code, 400)
        empty = load_workbook(BytesIO(self.client.get('/rtis/analysis.xlsx?division=MLDT').content))
        self.assertEqual(empty.active.max_row, 1)
        empty.close()

    def test_time_range_boundaries_and_excel_for_both_analyses(self):
        for analysis, train in [('passenger', '00441'), ('goods', 'AB123')]:
            event = 'J' if analysis == 'passenger' else 'K'
            for stamp in ['09:59:59', '10:00:00', '10:30:00', '11:00:00', '11:00:01', '23:59:59']:
                import_rtis(self.session, analysis + stamp.replace(':','') + '.xlsx',
                            workbook(train=train, event=event, speed='50', time='2026-09-14 ' + stamp), 'SDAH')
            query = f'analysis={analysis}&day=2026-09-14&time_from=10:00:00&time_to=11:00:00&event=JK&speed=40'
            response = self.client.get('/rtis?' + query)
            self.assertEqual(response.context['total'], 3)
            self.assertEqual(response.context['counts'][event], 3)
            self.assertIn('type="time" name="time_from"', response.text)
            for key in ('query', 'switch_query', 'division_query'):
                self.assertIn('time_from=10%3A00%3A00', response.context[key])
                self.assertIn('time_to=11%3A00%3A00', response.context[key])
            book = load_workbook(BytesIO(self.client.get('/rtis/analysis.xlsx?' + query).content))
            self.assertEqual(book.active.max_row, 4)
            self.assertEqual(book.active['A2'].value, datetime(2026,9,14,11))
            book.close()
            self.assertEqual(self.client.get(f'/rtis?analysis={analysis}&day=2026-09-14').context['total'], 6)
            self.assertEqual(self.client.get(f'/rtis?analysis={analysis}&day=2026-09-14&time_from=23:59:59').context['total'], 1)
            self.assertEqual(self.client.get(f'/rtis?analysis={analysis}&day=2026-09-14&time_to=10:00').context['total'], 2)
        for query in ['time_from=10:00', 'day=2026-09-14&time_from=bad',
                      'day=2026-09-14&time_to=25:00', 'day=2026-09-14&time_from=12:00&time_to=10:00']:
            for route in ['/rtis', '/rtis/analysis.xlsx']:
                self.assertEqual(self.client.get(route + '?' + query).status_code, 400)

    def test_train_search_filters_full_dataset_and_excel(self):
        for analysis, train, event, query in [('passenger','00441','J','0044'), ('goods','AB/123','K','ab/12')]:
            for i in range(101):
                import_rtis(self.session, f'{analysis}{i}.xlsx', workbook(
                    train=train, event=event, speed='50', time=f'2026-09-14 {i//60:02}:{i%60:02}:00'), 'SDAH')
            url = f'analysis={analysis}&train_no={query}&event=JK&speed=40&day=2026-09-14&time_from=00:00&time_to=02:00'
            page = self.client.get('/rtis?' + url + '&page=2')
            self.assertEqual(page.context['total'], 101)
            self.assertEqual(len(page.context['rows']), 1)
            self.assertEqual(page.context['counts'][event], 101)
            self.assertIn('name="train_no"', page.text)
            from urllib.parse import parse_qs
            for key in ('query','division_query','switch_query'):
                self.assertEqual(parse_qs(page.context[key])['train_no'], [query])
            book = load_workbook(BytesIO(self.client.get('/rtis/analysis.xlsx?' + url).content))
            self.assertEqual(book.active.max_row, 102)
            self.assertEqual({r[4] for r in book.active.iter_rows(min_row=2, values_only=True)}, {train})
            book.close()
            self.assertEqual(self.client.get(f'/rtis?analysis={analysis}&train_no=  {train}  ').context['total'], 101)
        for query in ['not-found', '%', '_']:
            response = self.client.get('/rtis', params={'train_no':query})
            self.assertEqual(response.context['total'], 0)
        self.assertEqual(self.client.get('/rtis', params={'train_no':' '}).context['total'], 101)
        self.assertEqual(self.client.get('/rtis', params={'train_no':'1'*101}).status_code, 400)
        station_page = self.client.get('/rtis', params={'station': 'bp'})
        self.assertEqual(station_page.context['total'], 101)
        self.assertEqual(station_page.context['station'], 'bp')
        self.assertIn('name="station"', station_page.text)
        self.assertEqual(self.client.get('/rtis', params={'station': 'not-found'}).context['total'], 0)
        self.assertEqual(self.client.get('/rtis', params={'station': '1' * 101}).status_code, 400)

    def test_all_divisions_analysis_filters_history_and_excel(self):
        divisions = {'SDAH','HWH','ASN','MLDT'}
        for division in sorted(divisions):
            for train, event in [('00441','J'), ('AB123','K'), ('','H')]:
                import_rtis(self.session, f'{division}{event}.xlsx', workbook(
                    division=division, train=train, event=event, speed='40'), division)
        for analysis, search, event in [('passenger','004','J'), ('goods','ab','K')]:
            query = f'division=ALL&analysis={analysis}&train_no={search}&event=JK&speed=40&day=2026-09-14&time_from=10:00&time_to=10:00'
            page = self.client.get('/rtis?' + query)
            self.assertEqual(page.context['total'], 4)
            self.assertEqual(page.context['counts'][event], 4)
            self.assertEqual({r.division for r in page.context['rows']}, divisions)
            self.assertEqual({r.division for r in page.context['history']}, divisions)
            self.assertEqual(len(page.context['history']), 12)
            self.assertIn('All divisions', page.text)
            self.assertIn('name="division" value="ALL"', page.text)
            self.assertIn('division=ALL', page.context['query'])
            self.assertIn('division=ALL', page.context['switch_query'])
            book = load_workbook(BytesIO(self.client.get('/rtis/analysis.xlsx?' + query).content))
            self.assertEqual(book.active.max_row, 5)
            self.assertEqual({r[1] for r in book.active.iter_rows(min_row=2, values_only=True)}, divisions)
            book.close()
            for division in divisions:
                single = self.client.get('/rtis?' + query.replace('division=ALL', 'division=' + division))
                self.assertEqual(single.context['total'], 1)
                self.assertEqual(len(single.context['history']), 3)
        self.assertEqual(self.client.get('/rtis?division=ALL&view=unclassified').context['total'], 4)
        self.assertEqual(self.client.post('/rtis/upload', data={'division':'ALL'},
                                         files={'files':('test.xlsx',workbook())}).status_code, 400)

    def test_empty_output_and_exports(self):
        response = self.client.get('/rtis/output')
        self.assertIn('No RTIS data uploaded yet', response.text)
        self.assertEqual(response.context['total'], 0)
        book = load_workbook(BytesIO(self.client.get('/rtis/output.xlsx').content))
        self.assertEqual(book.active.max_row, 1)
        book.close()
        pdf = PdfReader(BytesIO(self.client.get('/rtis/output.pdf').content))
        self.assertEqual(len(pdf.pages), 1)
        self.assertIn('No rows available.', ' '.join(pdf.pages[0].extract_text().split()))


if __name__ == '__main__':
    unittest.main()
