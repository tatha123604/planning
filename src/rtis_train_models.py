"""Saved passenger train shortlists from crew-link Excel workbooks."""
from __future__ import annotations
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
import re
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlmodel import Field, SQLModel, Session, select


class RtisTrainModel(SQLModel, table=True):
    id: int | None = Field(default=None, primary_key=True)
    filename: str
    digest: str = Field(unique=True)
    uploaded_at: datetime = Field(default_factory=datetime.utcnow)
    eligible_rows: int
    train_count: int
    mapping_json: str


def parse_train_model(content: bytes):
    if len(content) > 25 * 1024 * 1024:
        raise ValueError('Model file must be 25 MB or smaller.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError('Model workbook is too large when expanded.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, TypeError, ParseError) as exc:
        raise ValueError('Cannot read model. Upload a valid .xlsx workbook.') from exc
    mapping = {}
    eligible_rows = 0
    found = False
    normalize = lambda value: ' '.join(str(value or '').upper().split()).strip().rstrip('.')
    try:
        for sheet in book:
            columns = None
            for number, row in enumerate(sheet.iter_rows(), 1):
                labels = [normalize(cell.value) for cell in row]
                train_columns = [i for i, label in enumerate(labels) if label in ('TRAIN NO', 'TRAIN NUMBER')]
                if train_columns and 'HQ OF CREW' in labels and 'NAME OF TRAIN' in labels:
                    columns = (train_columns, labels.index('HQ OF CREW'), labels.index('NAME OF TRAIN'))
                    found = True
                    continue
                if columns is None:
                    continue
                trains, hq_column, name_column = columns
                hq = normalize(row[hq_column].value)
                if hq not in ('SDAH', 'KOAA'):
                    continue
                name = ' '.join(str(row[name_column].value or '').split())
                if not name:
                    raise ValueError(f'{sheet.title} row {number}: missing train name.')
                numbers = []
                for index in trains:
                    cell = row[index]
                    value = cell.value
                    if value is None or str(value).strip() == '':
                        continue
                    if isinstance(value, (float, int)) and not isinstance(value, bool) and float(value).is_integer():
                        train = str(int(value))
                        if re.fullmatch('0+', cell.number_format or ''):
                            train = train.zfill(len(cell.number_format))
                    else:
                        train = str(value).strip()
                    if not re.fullmatch(r'[0-9]+', train):
                        raise ValueError(f'{sheet.title} row {number}: invalid passenger train number {train!r}.')
                    numbers.append(train)
                if not numbers:
                    raise ValueError(f'{sheet.title} row {number}: missing train numbers.')
                eligible_rows += 1
                for train in numbers:
                    if train in mapping and mapping[train]['name'] != name:
                        raise ValueError(f'Conflicting train names for {train}. Correct the model and upload again.')
                    mapping[train] = {'name': name, 'hq': hq}
                if len(mapping) > 5000:
                    raise ValueError('Model may contain at most 5,000 train numbers.')
    except (BadZipFile, KeyError, OSError, ParseError) as exc:
        raise ValueError('Cannot read model. Upload a valid .xlsx workbook.') from exc
    finally:
        book.close()
    if not found:
        raise ValueError('Required headers: Train No, NAME OF TRAIN, HQ OF CREW.')
    if not mapping:
        raise ValueError('No train numbers with HQ OF CREW = SDAH or KOAA found.')
    return mapping, eligible_rows


def save_train_model(session: Session, filename: str, content: bytes):
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('Upload an .xlsx model workbook.')
    digest = sha256(content).hexdigest()
    existing = session.exec(select(RtisTrainModel).where(RtisTrainModel.digest == digest)).first()
    if existing:
        return existing
    mapping, eligible_rows = parse_train_model(content)
    model = RtisTrainModel(filename=filename.replace('\\', '/').split('/')[-1], digest=digest,
                          eligible_rows=eligible_rows, train_count=len(mapping), mapping_json=json.dumps(mapping))
    try:
        session.add(model)
        session.commit()
        session.refresh(model)
    except IntegrityError:
        session.rollback()
        return session.exec(select(RtisTrainModel).where(RtisTrainModel.digest == digest)).one()
    return model


def selected_train_model(session: Session, model_id: int, analysis: str):
    if not model_id:
        return None, {}
    if analysis != 'passenger':
        raise HTTPException(400, 'Train number models apply to Passenger analysis only.')
    model = session.get(RtisTrainModel, model_id)
    if model is None:
        raise HTTPException(404, 'Train number model not found. Select an available model.')
    return model, json.loads(model.mapping_json)
