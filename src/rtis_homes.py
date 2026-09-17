"""FSD home-signal mappings used by RTIS J/K events."""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
import math
import re

from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlmodel import Field, Session, SQLModel, select


class RtisHomeModel(SQLModel, table=True):
    id: int | None = Field(default=None, primary_key=True)
    filename: str
    digest: str = Field(unique=True)
    uploaded_at: datetime = Field(default_factory=datetime.utcnow)
    signal_count: int
    station_count: int
    mapped_station_count: int
    mapping_json: str


def _number(value, label):
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label} must be a number.") from None
    if not math.isfinite(result):
        raise ValueError(f"{label} must be a finite number.")
    return result


def parse_home_model(content: bytes):
    if len(content) > 25 * 1024 * 1024:
        raise ValueError("FSD file must be 25 MB or smaller.")
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Cannot read FSD file. Upload a valid .xlsx workbook.") from exc
    mapping = {}
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            labels = [' '.join(str(value or '').upper().split()) for value in header]
            required = ('STATION', 'DIRN', 'TYPE', 'LATITUDE', 'LONGITUDE')
            if not all(name in labels for name in required):
                continue
            indexes = {name: labels.index(name) for name in required}
            for row_number, row in enumerate(rows, 2):
                station = str(row[indexes['STATION']] or '').strip().upper()
                direction = str(row[indexes['DIRN']] or '').strip().upper()
                signal_type = str(row[indexes['TYPE']] or '').strip()
                if not station or not direction or signal_type.upper() not in ('HOME', 'I/HOME'):
                    continue
                if direction.startswith('UP'):
                    event = 'J'
                    label = 'UP HOME'
                elif direction.startswith('DN'):
                    event = 'K'
                    label = 'DOWN HOME'
                else:
                    continue
                latitude = _number(row[indexes['LATITUDE']], f"{sheet.title} row {row_number} Latitude")
                longitude = _number(row[indexes['LONGITUDE']], f"{sheet.title} row {row_number} Longitude")
                if not -90 <= latitude <= 90 or not -180 <= longitude <= 180:
                    raise ValueError(f"{sheet.title} row {row_number}: coordinates are out of range.")
                key = f"{station}|{event}"
                mapping.setdefault(key, {'station': station, 'event': event, 'label': label, 'homes': []})['homes'].append({
                    'line': direction, 'type': signal_type, 'latitude': latitude, 'longitude': longitude,
                })
    finally:
        book.close()
    if not mapping:
        raise ValueError("No Home or I/Home rows with UP/DN direction found.")
    return mapping, sum(len(value['homes']) for value in mapping.values())


def _haversine_m(latitude1, longitude1, latitude2, longitude2):
    radius = 6371008.8
    lat1, lat2 = math.radians(latitude1), math.radians(latitude2)
    dlat = lat2 - lat1
    dlon = math.radians(longitude2 - longitude1)
    value = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return round(2 * radius * math.asin(math.sqrt(min(1, value))), 1)


def _polygon_centroid(polygon):
    """Return the area-weighted centroid of (longitude, latitude) points."""
    area_twice = 0.0
    longitude_sum = 0.0
    latitude_sum = 0.0
    for first, second in zip(polygon, polygon[1:] + polygon[:1]):
        cross = first[0] * second[1] - second[0] * first[1]
        area_twice += cross
        longitude_sum += (first[0] + second[0]) * cross
        latitude_sum += (first[1] + second[1]) * cross
    if abs(area_twice) < 1e-12:
        return (sum(point[0] for point in polygon) / len(polygon),
                sum(point[1] for point in polygon) / len(polygon))
    return longitude_sum / (3 * area_twice), latitude_sum / (3 * area_twice)


def enrich_home_model(mapping, geofence_polygons):
    mapped = 0
    for value in mapping.values():
        polygon = geofence_polygons.get(value['station'], [])
        if polygon:
            station_lon, station_lat = _polygon_centroid(polygon)
            value['station_latitude'] = station_lat
            value['station_longitude'] = station_lon
            for home in value['homes']:
                home['distance_m'] = _haversine_m(station_lat, station_lon, home['latitude'], home['longitude'])
            mapped += 1
        else:
            for home in value['homes']:
                home['distance_m'] = None
    return mapped


def save_home_model(session: Session, filename: str, content: bytes, geofence_polygons=None):
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('Upload an .xlsx FSD workbook.')
    digest = sha256(content).hexdigest()
    existing = session.exec(select(RtisHomeModel).where(RtisHomeModel.digest == digest)).first()
    if existing:
        if geofence_polygons:
            mapping = json.loads(existing.mapping_json)
            mapped = enrich_home_model(mapping, geofence_polygons)
            existing.mapping_json = json.dumps(mapping)
            existing.mapped_station_count = mapped
            session.add(existing)
            session.commit()
            session.refresh(existing)
        return existing
    mapping, signal_count = parse_home_model(content)
    mapped = enrich_home_model(mapping, geofence_polygons or {})
    model = RtisHomeModel(filename=filename.replace('\\', '/').split('/')[-1], digest=digest,
                           signal_count=signal_count, station_count=len({v['station'] for v in mapping.values()}),
                           mapped_station_count=mapped, mapping_json=json.dumps(mapping))
    try:
        session.add(model)
        session.commit()
        session.refresh(model)
    except IntegrityError:
        session.rollback()
        return session.exec(select(RtisHomeModel).where(RtisHomeModel.digest == digest)).one()
    return model


def selected_home_model(session: Session):
    model = session.exec(select(RtisHomeModel).order_by(RtisHomeModel.id.desc())).first()
    return (model, json.loads(model.mapping_json)) if model else (None, {})


def event_home_details(train_station, event_type, mapping):
    if event_type not in ('J', 'K'):
        return '', ''
    value = mapping.get(f"{str(train_station or '').strip().upper()}|{event_type}")
    if not value:
        return f"{'UP' if event_type == 'J' else 'DOWN'} HOME", ''
    distances = [home['distance_m'] for home in value['homes'] if home.get('distance_m') is not None]
    distance = min(distances) if distances else ''
    return value['label'], distance
