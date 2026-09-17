"""Exports for the complete RTIS table, using the project's report styling."""
from datetime import date, datetime
from io import BytesIO
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .table_pdf import _build_table_pdf_bytes


def _excel_value(header: str, value: str):
    if not value:
        return None, "General"
    try:
        if header == "Sr.No.":
            return int(value), "0"
        if header in ("Latitude", "Longitude", "Speed", "Home Distance (m)", "Linear distance (m)"):
            return float(value), "0.######"
        if header in ("Event Time", "Event date_time", "Reporting Time"):
            return datetime.fromisoformat(value), "yyyy-mm-dd hh:mm:ss"
        if header == "Train Start Date":
            return date.fromisoformat(value), "yyyy-mm-dd"
    except (ValueError, OverflowError):
        pass
    # Identifiers and unrecognized source values remain literal text.
    return value, "@"


def build_rtis_excel(headers: Sequence[str], rows: Iterable[list[str]]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("RTIS Output")
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    widths = {"Sr.No.": 10, "Device Id": 14, "Loco No.": 14, "Latitude": 14,
              "Longitude": 14, "Station": 12, "Event Time": 23, "Event date_time": 23,
              "Event Type": 12, "Speed": 10, "Division Code": 15, "Reporting Time": 23,
              "Train Number": 32, "Train Start Date": 18, "Train Name": 38, "HQ OF CREW": 16,
              "Home Signal": 16, "Home Distance (m)": 20, "Linear distance (m)": 20, "Direction": 12,
              "Type": 14, "Station DIRN": 24}
    for index, header in enumerate(headers, start=1):
        width = widths.get(header, 18)
        sheet.column_dimensions[get_column_letter(index)].width = width
    header_font = Font(name="Arial", size=10, bold=True, color="F3FBFF")
    body_font = Font(name="Arial", size=10, color="16314C")
    header_fill = PatternFill(fill_type="solid", fgColor="16314C")
    header_alignment = Alignment(horizontal="center", vertical="center")
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font, cell.fill, cell.alignment = header_font, header_fill, header_alignment
        header_cells.append(cell)
    sheet.row_dimensions[1].height = 24
    sheet.append(header_cells)
    row_count = 0
    for row in rows:
        cells = []
        for header, raw in zip(headers, row):
            value, number_format = _excel_value(header, raw)
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.number_format = number_format
            cell.font = body_font
            cells.append(cell)
        sheet.append(cells)
        row_count += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_count + 1}"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_rtis_pdf(headers: Sequence[str], rows: Iterable[list[str]]) -> bytes:
    return _build_table_pdf_bytes(
        "RTIS Output - All Divisions", list(headers), list(rows),
        "Divisions: SDAH, HWH, ASN, MLDT",
        page_width=1191.0, page_height=842.0, body_font_size=8.5, compress=True,
    )
