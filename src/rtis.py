"""Division-wise RTIS imports and event browsing."""
from __future__ import annotations

from collections import Counter
import csv
from datetime import date, datetime, timedelta, timezone, time as clock_time
from zoneinfo import ZoneInfo
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
import posixpath
import re
from urllib.parse import urlencode
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile
from openpyxl import load_workbook

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import RedirectResponse, Response
from fastapi.templating import Jinja2Templates
from sqlalchemy import Column, LargeBinary, UniqueConstraint, func
from sqlalchemy.exc import IntegrityError
from sqlmodel import Field, SQLModel, Session, select

from .db import get_session
from .rtis_exports import build_rtis_excel, build_rtis_pdf
from .rtis_train_models import RtisTrainModel, save_train_model, selected_train_model
from .rtis_homes import RtisHomeModel, event_home_details, save_home_model, selected_home_model

DIVISIONS = ("SDAH", "HWH", "ASN", "MLDT")
FOCUS_EVENTS = ("H", "J", "K")
ANALYSIS_TYPES = {"passenger": "Passenger Train Analysis", "goods": "Goods Train Analysis"}
MAX_UPLOAD_MB = 25
MAX_BYTES = MAX_UPLOAD_MB * 1024 * 1024
RTIS_RETENTION_DAYS = 2
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
HEADERS = ("Sr.No.", "Device Id", "Loco No.", "Latitude", "Longitude", "Station",
           "Event Time", "Event Type", "Speed", "Division Code", "Reporting Time",
           "Train Number", "Train Start Date")
router = APIRouter()
templates = Jinja2Templates(directory=str(Path(__file__).resolve().parent.parent / "templates"))
IST = ZoneInfo("Asia/Kolkata")
def format_ist(value):
    if not value:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(IST).strftime("%d-%m-%Y %H:%M")
templates.env.filters["ist"] = format_ist


class RtisUpload(SQLModel, table=True):
    __table_args__ = (UniqueConstraint("division", "digest"),)
    id: int | None = Field(default=None, primary_key=True)
    division: str = Field(index=True)
    filename: str
    digest: str
    uploaded_at: datetime = Field(default_factory=datetime.utcnow)
    first_day: date
    last_day: date
    row_count: int
    added_count: int
    event_counts: str
    content: bytes = Field(sa_column=Column(LargeBinary, nullable=False))


class RtisEvent(SQLModel, table=True):
    id: int | None = Field(default=None, primary_key=True)
    upload_id: int = Field(foreign_key="rtisupload.id", index=True)
    fingerprint: str = Field(unique=True, index=True)
    division: str = Field(index=True)
    event_type: str = Field(index=True)
    event_time: datetime = Field(index=True)
    station: str = Field(index=True)
    loco: str
    train: str
    speed: float | None = None
    latitude: float | None = None
    longitude: float | None = None
    source: str


class RtisAnalysisUpload(SQLModel, table=True):
    """Uploaded secondary GPS data and the filters used for its analysis."""
    id: int | None = Field(default=None, primary_key=True)
    filename: str
    uploaded_at: datetime = Field(default_factory=datetime.utcnow)
    analysis_date: str
    train_type: str
    train_no: str
    loco_no: str
    time_from: str
    time_to: str
    content: bytes = Field(sa_column=Column(LargeBinary, nullable=False))
    station_from: str = ""
    station_to: str = ""


class RtisAnalysisPrimaryUpload(SQLModel, table=True):
    id: int | None = Field(default=None, primary_key=True)
    analysis_upload_id: int = Field(index=True)
    filename: str
    content: bytes = Field(sa_column=Column(LargeBinary, nullable=False))


def train_type_filters():
    """User rule: digits only = passenger; letters plus digits = goods."""
    train = func.trim(RtisEvent.train)
    passenger = (train != "") & ~train.op("GLOB")("*[^0-9]*")
    goods = train.op("GLOB")("*[A-Za-z]*") & train.op("GLOB")("*[0-9]*")
    return passenger, goods


def _xml(archive: ZipFile, name: str):
    data = archive.read(name)
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValueError("Unsupported XML declarations in workbook.")
    return ET.fromstring(data)


def _timestamp(value: str, epoch: datetime) -> datetime:
    try:
        result = datetime.fromisoformat(value.strip())
        if result.tzinfo is not None:
            raise ValueError("Use local event times without timezone offsets.")
        return result
    except ValueError:
        try:
            serial = float(value)
            # Excel's Windows date system contains a fictitious 1900 leap day.
            adjustment = 1 if epoch.year == 1899 and 0 < serial < 60 else 0
            return epoch + timedelta(days=serial + adjustment)
        except (ValueError, OverflowError):
            raise ValueError("Expected an ISO date/time or Excel date value.") from None


def parse_rtis(content: bytes, division: str) -> list[dict]:
    """Read cell values directly: RTIS exports may contain invalid Excel fills.

    Formula cells are rejected, not evaluated. Original bytes are kept unchanged.
    """
    if division not in DIVISIONS:
        raise ValueError("Choose SDAH, HWH, ASN or MLDT.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError("Workbook expands beyond the 150 MB limit.")
            shared = []
            if "xl/sharedStrings.xml" in archive.namelist():
                shared = ["".join(node.itertext()) for node in _xml(archive, "xl/sharedStrings.xml")]
            workbook = _xml(archive, "xl/workbook.xml")
            properties = workbook.find("s:workbookPr", NS)
            epoch = datetime(1904, 1, 1) if properties is not None and properties.get("date1904") in ("1", "true") else datetime(1899, 12, 30)
            links = {node.get("Id"): node for node in _xml(archive, "xl/_rels/workbook.xml.rels")}
            records = []
            for sheet in workbook.findall("s:sheets/s:sheet", NS):
                link = links[sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")]
                if link.get("TargetMode") == "External":
                    raise ValueError("External worksheet links are unsupported.")
                target = link.get("Target", "")
                target = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
                if not target.startswith("xl/worksheets/"):
                    continue
                header = None
                for row in _xml(archive, target).findall("s:sheetData/s:row", NS):
                    cells = {}
                    for cell in row:
                        column = "".join(c for c in cell.get("r", "") if c.isalpha())
                        if cell.find("s:f", NS) is not None:
                            raise ValueError(f"{sheet.get('name')} row {row.get('r')}: formulas are not accepted in RTIS data.")
                        value = cell.find("s:v", NS)
                        raw = value.text or "" if value is not None else ""
                        if cell.get("t") == "s":
                            raw = shared[int(raw)]
                        elif cell.get("t") == "inlineStr":
                            raw = "".join(cell.find("s:is", NS).itertext())
                        cells[column] = raw.strip()
                    if not any(cells.values()):
                        continue
                    if header is None:
                        header = cells
                        missing = set(HEADERS) - set(header.values())
                        if missing:
                            raise ValueError(f"{sheet.get('name')}: missing columns: {', '.join(sorted(missing))}.")
                        if len(set(header.values())) != len(header.values()):
                            raise ValueError("Duplicate column names in workbook.")
                        continue
                    source = {name: cells.get(col, "") for col, name in header.items() if name in HEADERS}
                    label = f"{sheet.get('name')} row {row.get('r')}"
                    if source["Division Code"].upper() != division:
                        raise ValueError(f"{label}: Division Code is '{source['Division Code']}', expected {division}.")
                    try:
                        event_time = _timestamp(source["Event Time"], epoch)
                        if source["Reporting Time"]:
                            _timestamp(source["Reporting Time"], epoch)
                        speed = float(source["Speed"]) if source["Speed"] else None
                        latitude = float(source["Latitude"]) if source["Latitude"] else None
                        longitude = float(source["Longitude"]) if source["Longitude"] else None
                        if latitude is not None and not -90 <= latitude <= 90:
                            raise ValueError("Latitude is outside the valid range.")
                        if longitude is not None and not -180 <= longitude <= 180:
                            raise ValueError("Longitude is outside the valid range.")
                        if speed is not None and (not math.isfinite(speed) or speed < 0):
                            raise ValueError("Speed must be a non-negative number.")
                    except ValueError as exc:
                        raise ValueError(f"{label}: {exc}") from None
                    event_type = source["Event Type"].upper()
                    if not event_type or not source["Loco No."]:
                        raise ValueError(f"{label}: Event Type and Loco No. are required.")
                    # Ignore export serial/report-delivery time when identifying the same event.
                    identity = [division, source["Device Id"], source["Loco No."], source["Station"],
                                event_time.isoformat(), event_type, source["Latitude"], source["Longitude"]]
                    records.append(dict(division=division, event_type=event_type, event_time=event_time,
                                        station=source["Station"], loco=source["Loco No."],
                                        train=source["Train Number"], speed=speed, latitude=latitude, longitude=longitude,
                                        source=json.dumps(source),
                                        fingerprint=sha256(json.dumps(identity).encode()).hexdigest()))
                    if len(records) > 100_000:
                        raise ValueError("Upload at most 100,000 events per file.")
            if not records:
                raise ValueError("No RTIS event rows found in the workbook.")
            return records
    except (BadZipFile, ET.ParseError, KeyError, IndexError, TypeError, AttributeError, RuntimeError, NotImplementedError):
        raise ValueError("Cannot read this workbook. Upload a valid RTIS .xlsx export.") from None


def import_rtis(session: Session, filename: str, content: bytes, division: str) -> str:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Upload an .xlsx file.")
    if len(content) > MAX_BYTES:
        raise ValueError(f"Each file must be {MAX_UPLOAD_MB} MB or smaller.")
    if division not in DIVISIONS:
        raise ValueError("Choose SDAH, HWH, ASN or MLDT.")
    digest = sha256(content).hexdigest()
    previous_upload = session.exec(select(RtisUpload).where(RtisUpload.division == division,
                                                              RtisUpload.digest == digest)).first()
    if previous_upload:
        # A previous upload can be left behind without events after an undo,
        # retention cleanup, or an interrupted duplicate import. In that case
        # allow the same workbook to restore its missing events.
        previous_event = session.exec(select(RtisEvent.id).where(RtisEvent.upload_id == previous_upload.id).limit(1)).first()
        if previous_event is not None:
            return "Already uploaded; no duplicate events added."
        session.delete(previous_upload)
        session.flush()
    records = parse_rtis(content, division)
    unique = {row["fingerprint"]: row for row in records}
    existing = set()
    keys = list(unique)
    for offset in range(0, len(keys), 500):
        existing.update(session.exec(select(RtisEvent.fingerprint).where(RtisEvent.fingerprint.in_(keys[offset:offset + 500]))).all())
    new = [row for key, row in unique.items() if key not in existing]
    upload = RtisUpload(division=division, filename=filename.replace("\\", "/").split("/")[-1], digest=digest,
                        first_day=min(row["event_time"].date() for row in records),
                        last_day=max(row["event_time"].date() for row in records),
                        row_count=len(records), added_count=len(new),
                        event_counts=json.dumps(Counter(row["event_type"] for row in records)), content=content)
    try:
        session.add(upload)
        session.flush()
        session.add_all([RtisEvent(upload_id=upload.id, **row) for row in new])
        session.commit()
    except IntegrityError:
        session.rollback()
        raise ValueError("An overlapping upload was saved at the same time. Retry this file.") from None
    prune_rtis_history(session)
    return f"Saved {len(new):,} new events; {len(records) - len(new):,} duplicate rows skipped."


def prune_rtis_history(session: Session) -> None:
    """Keep RTIS events for the latest two event dates in the database."""
    latest = session.exec(select(func.max(RtisEvent.event_time))).one()
    if latest is None:
        return
    cutoff = latest.date() - timedelta(days=RTIS_RETENTION_DAYS - 1)
    old_events = session.exec(select(RtisEvent).where(RtisEvent.event_time < datetime.combine(cutoff, clock_time.min))).all()
    if not old_events:
        return
    old_upload_ids = {event.upload_id for event in old_events}
    for event in old_events:
        session.delete(event)
    session.flush()
    for upload_id in old_upload_ids:
        if session.exec(select(RtisEvent.id).where(RtisEvent.upload_id == upload_id)).first() is None:
            upload = session.get(RtisUpload, upload_id)
            if upload:
                session.delete(upload)
    session.commit()


def analysis_filters(division, day, event, analysis, view, speed, time_from="", time_to="", train_no="", station="", model_trains=None, train_name="", loco_no=""):
    """Keep on-screen results and full Excel downloads on the same filter rules."""
    if speed not in ("", "30", "40", "50"):
        raise HTTPException(400, "Choose All speeds, 30+, 40+ or 50+.")
    if analysis not in ANALYSIS_TYPES or view not in ("classified", "unclassified"):
        raise HTTPException(400, "Invalid analysis selection.")
    if division not in (*DIVISIONS, "ALL") or event not in (*FOCUS_EVENTS, "HJK", "JK"):
        raise HTTPException(400, "Invalid division or event filter.")
    try:
        selected_day = date.fromisoformat(day) if day else None
    except ValueError:
        raise HTTPException(400, "Invalid date.") from None
    if (time_from or time_to) and not selected_day:
        raise HTTPException(400, "Select an event date when filtering by time.")
    parsed_times = []
    for value in (time_from, time_to):
        try:
            if value and (len(value) not in (5, 8) or value[2] != ':' or
                          (len(value) == 8 and value[5] != ':')):
                raise ValueError
            parsed_times.append(clock_time.fromisoformat(value) if value else None)
        except ValueError:
            raise HTTPException(400, "Enter times as HH:MM or HH:MM:SS.") from None
    start_time, end_time = parsed_times
    if start_time and end_time and start_time > end_time:
        raise HTTPException(400, "To time must be on or after From time within the selected date.")
    filters = [RtisEvent.event_type.in_(FOCUS_EVENTS)]
    if model_trains is not None:
        filters.append(RtisEvent.train.in_(list(model_trains)))
    train_name = train_name.strip()
    if train_name:
        if analysis != "passenger" or model_trains is None:
            raise HTTPException(400, "Select a Passenger train no. model before filtering by Train Name.")
        matching_trains = [train for train, value in model_trains.items()
                           if str(value.get("name", "")).casefold() == train_name.casefold()]
        filters.append(RtisEvent.train.in_(matching_trains or ["__no_train_name_match__"]))
    filters.append(RtisEvent.division.in_(DIVISIONS) if division == "ALL" else RtisEvent.division == division)
    train_no = train_no.strip()
    if len(train_no) > 100:
        raise HTTPException(400, "Train no. search must be 100 characters or fewer.")
    if train_no:
        filters.append(func.lower(RtisEvent.train).contains(train_no.lower(), autoescape=True))
    station = station.strip()
    if len(station) > 100:
        raise HTTPException(400, "Station search must be 100 characters or fewer.")
    if station:
        filters.append(func.lower(RtisEvent.station).contains(station.lower(), autoescape=True))
    loco_no = loco_no.strip()
    if len(loco_no) > 100:
        raise HTTPException(400, "Loco no. search must be 100 characters or fewer.")
    if loco_no:
        filters.append(func.lower(RtisEvent.loco).contains(loco_no.lower(), autoescape=True))
    if speed:
        filters.append(RtisEvent.speed >= int(speed))
    if selected_day:
        try:
            start = datetime.combine(selected_day, start_time or clock_time.min)
            # Include the entire selected end second; blank end covers the full day.
            end = (datetime.combine(selected_day, end_time) + timedelta(seconds=1)
                   if end_time else datetime.combine(selected_day + timedelta(days=1), clock_time.min))
        except OverflowError:
            raise HTTPException(400, "Event date is out of range.") from None
        filters.extend([RtisEvent.event_time >= start, RtisEvent.event_time < end])
    passenger, goods = train_type_filters()
    unknown = ~(passenger | goods)
    unknown_filters = [*filters, unknown]
    if view == "unclassified":
        filters.append(unknown)
    else:
        filters.append(passenger if analysis == "passenger" else goods)
    summary_filters = list(filters)
    if event == "JK":
        filters.append(RtisEvent.event_type.in_(("J", "K")))
    elif event != "HJK":
        filters.append(RtisEvent.event_type == event)
    return filters, summary_filters, unknown_filters


def rtis_ordering(sort_by="event_time", sort_order="desc"):
    fields = {
        "event_time": RtisEvent.event_time,
        "division": RtisEvent.division,
        "station": RtisEvent.station,
        "event_type": RtisEvent.event_type,
        "train": RtisEvent.train,
        "loco": RtisEvent.loco,
        "speed": RtisEvent.speed,
    }
    if sort_by not in fields or sort_order not in ("asc", "desc"):
        raise HTTPException(400, "Invalid table sort selection.")
    field = fields[sort_by]
    ordered = field.asc() if sort_order == "asc" else field.desc()
    return [ordered.nulls_last(), RtisEvent.id.desc()]


@router.get("/rtis")
def rtis_page(request: Request, division: str = "ALL", day: str = "", event: str = "HJK",
              page: int = 1, analysis: str = "passenger", view: str = "classified",
              speed: str = "", time_from: str = "", time_to: str = "", train_no: str = "", station: str = "", home_station: str = "", train_name: str = "", loco_no: str = "", sort_by: str = "event_time", sort_order: str = "desc",
              model_id: int = 0,
              session: Session = Depends(get_session)):
    train_no = train_no.strip()
    station = station.strip()
    home_station = home_station.strip().upper()
    train_name = train_name.strip()
    loco_no = loco_no.strip()
    if len(home_station) > 100:
        raise HTTPException(400, "FSD station search must be 100 characters or fewer.")
    ordering = rtis_ordering(sort_by, sort_order)
    selected_model, train_names = selected_train_model(session, model_id, analysis)
    home_model, home_mapping = selected_home_model(session)
    home_signal_rows = []
    for value in home_mapping.values():
        for home in value.get("homes", []):
            home_signal_rows.append({
                "station": value["station"],
                "direction": "UP" if value["event"] == "J" else "DOWN",
                "event": value["event"],
                "type": home.get("type", ""),
                "station_latitude": value.get("station_latitude"),
                "station_longitude": value.get("station_longitude"),
                "latitude": home.get("latitude"),
                "longitude": home.get("longitude"),
                "station_dirn": home.get("line", ""),
                "distance_m": home.get("distance_m"),
            })
    home_signal_rows.sort(key=lambda row: (row["station"], row["direction"], row["station_dirn"], row["type"], row["latitude"], row["longitude"]))
    home_signal_match_count = sum(not home_station or home_station in row["station"] for row in home_signal_rows)
    filters, summary_filters, unknown_filters = analysis_filters(division, day, event, analysis, view, speed, time_from, time_to, train_no, station, train_names if selected_model else None, train_name, loco_no)
    unclassified = session.exec(select(func.count()).select_from(RtisEvent).where(*unknown_filters)).one()
    counts = dict(session.exec(select(RtisEvent.event_type, func.count()).where(*summary_filters)
                              .group_by(RtisEvent.event_type)).all())
    total = session.exec(select(func.count()).select_from(RtisEvent).where(*filters)).one()
    pages = max(1, (total + 99) // 100)
    page = min(max(1, page), pages)
    rows = session.exec(select(RtisEvent).where(*filters).order_by(*ordering).offset((page - 1) * 100).limit(100)).all()
    # Select metadata only, never all the stored workbooks while rendering the page.
    history_filter = RtisUpload.division.in_(DIVISIONS) if division == "ALL" else RtisUpload.division == division
    history = session.exec(select(RtisUpload.id, RtisUpload.filename, RtisUpload.uploaded_at, RtisUpload.division,
                                  RtisUpload.first_day, RtisUpload.last_day, RtisUpload.row_count,
                                  RtisUpload.added_count).where(history_filter)
                           .order_by(RtisUpload.id.desc()).limit(50)).all()
    return templates.TemplateResponse(request=request, name="rtis.html", context={
        "active_page": "rtis", "divisions": DIVISIONS, "division": division, "day": day,
        "division_label": "All divisions" if division == "ALL" else division,
        "max_upload_mb": MAX_UPLOAD_MB,
        "event": event, "counts": counts, "rows": rows, "total": total, "history": history,
        "page": page, "pages": pages,
        "model_id": model_id, "selected_model": selected_model, "train_names": train_names,
        "home_model": home_model, "home_details": {row.id: event_home_details(row.station, row.event_type, home_mapping) for row in rows},
        "home_signal_rows": home_signal_rows,
        "home_signal_match_count": home_signal_match_count,
        "train_models": session.exec(select(RtisTrainModel.id, RtisTrainModel.filename, RtisTrainModel.train_count,
                                            RtisTrainModel.eligible_rows).order_by(RtisTrainModel.id.desc())).all(),
        "analysis": analysis, "analysis_types": ANALYSIS_TYPES, "view": view,
        "unclassified": unclassified, "speed": speed, "time_from": time_from, "time_to": time_to, "train_no": train_no, "station": station,
        "home_station": home_station,
        "train_name": train_name, "loco_no": loco_no,
        "sort_by": sort_by, "sort_order": sort_order,
        "train_name_options": sorted({str(value.get("name", "")).strip() for value in train_names.values() if value.get("name")}),
        "division_query": urlencode({"day": day, "event": event, "analysis": analysis, "view": view, "speed": speed, "time_from": time_from, "time_to": time_to, "train_no": train_no, "station": station, "train_name": train_name, "loco_no": loco_no, "sort_by": sort_by, "sort_order": sort_order, "model_id": model_id}),
        "switch_query": urlencode({"division": division, "day": day, "event": event, "speed": speed, "time_from": time_from, "time_to": time_to, "train_no": train_no, "station": station, "sort_by": sort_by, "sort_order": sort_order}),
        "query": urlencode({"division": division, "day": day, "event": event, "analysis": analysis, "view": view, "speed": speed, "time_from": time_from, "time_to": time_to, "train_no": train_no, "station": station, "train_name": train_name, "loco_no": loco_no, "sort_by": sort_by, "sort_order": sort_order, "model_id": model_id}),
        "notice": request.query_params.get("notice", ""),
    })


@router.get("/rtis/analysis.xlsx")
def rtis_analysis_excel(division: str = "ALL", day: str = "", event: str = "HJK",
                        analysis: str = "passenger", view: str = "classified", speed: str = "",
                        time_from: str = "", time_to: str = "", train_no: str = "", station: str = "", train_name: str = "", loco_no: str = "", sort_by: str = "event_time", sort_order: str = "desc",
                        model_id: int = 0,
                        session: Session = Depends(get_session)):
    selected_model, train_names = selected_train_model(session, model_id, analysis)
    home_model, home_mapping = selected_home_model(session)
    ordering = rtis_ordering(sort_by, sort_order)
    filters, _, _ = analysis_filters(division, day, event, analysis, view, speed, time_from, time_to, train_no, station, train_names if selected_model else None, train_name, loco_no)
    events = session.exec(select(RtisEvent).where(*filters)
                          .order_by(*ordering)
                          .execution_options(yield_per=1000))
    headers = ("Event date_time", "Division Code", "Station", "Event Type", "Train Number", "Loco No.", "Speed")
    if selected_model:
        headers = (*headers, "Train Name", "HQ OF CREW")
    if home_model:
        headers = (*headers, "Home Signal", "Home Distance (m)")
    rows = ([row.event_time.isoformat(sep=" "), row.division, row.station, row.event_type,
             row.train, row.loco, str(row.speed) if row.speed is not None else ""] +
            ([train_names[row.train]['name'], train_names[row.train]['hq']] if selected_model else []) +
            ([*event_home_details(row.station, row.event_type, home_mapping)] if home_model else []) for row in events)
    period = f"_{time_from.replace(':', '') or '000000'}-{time_to.replace(':', '') or '235959'}" if time_from or time_to else ""
    filename = f"RTIS_{division}_{analysis}_{view}_{day or 'all-dates'}{period}_{event}_{speed or 'all'}-speed.xlsx"
    if selected_model:
        filename = filename.replace('.xlsx', f'_model-{model_id}.xlsx')
    return Response(content=build_rtis_excel(headers, rows),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _rtis_ssts_geofences():
    try:
        from .app import _fetch_ssts_geofence_polygons, fetch_ssts_token
        return _fetch_ssts_geofence_polygons(fetch_ssts_token())
    except Exception:
        return {}


@router.get("/rtis/event-analysis/{event_id}")
def rtis_event_analysis(event_id: int, request: Request, session: Session = Depends(get_session)):
    def error_page(message: str, status_code: int = 400):
        response = templates.TemplateResponse(request=request, name="rtis_analysis_error.html", context={"request": request, "active_page": "rtis", "message": message})
        response.status_code = status_code
        return response
    event = session.get(RtisEvent, event_id)
    if event is None:
        return error_page("RTIS event not found.", 404)
    if event.latitude is None or event.longitude is None:
        return error_page("This RTIS event has no latitude/longitude in the uploaded file.")
    home_model, mapping = selected_home_model(session)
    if not home_model:
        return error_page("Upload an FSD signal model before running analysis.")
    candidates = []
    station_code = str(event.station or "").strip().upper()
    for value in mapping.values():
        if str(value.get("station") or "").strip().upper() != station_code:
            continue
        if value.get("event") != event.event_type:
            continue
        for home in value.get("homes", []):
            if home.get("latitude") is not None and home.get("longitude") is not None:
                distance = (float(home["latitude"]) - event.latitude) ** 2 + (float(home["longitude"]) - event.longitude) ** 2
                candidates.append((distance, value, home))
    # H is an approach event, so use the nearest FSD home at the same station
    # when there is no direction-specific J/K signal mapping.
    if not candidates and event.event_type == "H":
        for value in mapping.values():
            if str(value.get("station") or "").strip().upper() != station_code:
                continue
            for home in value.get("homes", []):
                if home.get("latitude") is not None and home.get("longitude") is not None:
                    distance = (float(home["latitude"]) - event.latitude) ** 2 + (float(home["longitude"]) - event.longitude) ** 2
                    candidates.append((distance, value, home))
    if not candidates:
        return error_page("No matching FSD home signal was found for this RTIS event. Check that the FSD model contains the same station and event direction.", 404)
    _, signal_group, home = min(candidates, key=lambda item: item[0])
    day_start = datetime.combine(event.event_time.date(), clock_time.min)
    day_end = day_start + timedelta(days=1)
    same_route_query = select(RtisEvent).where(
        RtisEvent.division == event.division,
        RtisEvent.event_time >= day_start,
        RtisEvent.event_time < day_end,
    )
    if event.train.strip():
        same_route_query = same_route_query.where(RtisEvent.train == event.train)
    else:
        same_route_query = same_route_query.where(RtisEvent.loco == event.loco)
    route_events = session.exec(same_route_query.order_by(RtisEvent.event_time)).all()
    route_points = [
        {"lat": row.latitude, "lon": row.longitude, "speed": row.speed or 0, "time": row.event_time.isoformat(), "station": row.station}
        for row in route_events if row.latitude is not None and row.longitude is not None
    ]
    event_point = {"lat": event.latitude, "lon": event.longitude, "speed": event.speed or 0, "time": event.event_time.isoformat(), "station": event.station}
    if not route_points:
        route_points = [event_point]
    signal = {"station": signal_group.get("station", ""), "event": event.event_type, "label": signal_group.get("label", "FSD Home Signal"), "line": home.get("line", ""), "lat": home["latitude"], "lon": home["longitude"], "speed": event.speed or 0, "time": event.event_time.isoformat(), "active": True}
    upload = {"filename": f"RTIS event {event.id}", "analysis_date": event.event_time.strftime("%Y-%m-%d"), "train_type": "RTIS event"}
    return templates.TemplateResponse(request=request, name="rtis_analysis_result.html", context={"request": request, "active_page": "rtis_analysis", "upload": upload, "points": route_points, "signals": [signal], "event_markers": [{"lat": event.latitude, "lon": event.longitude, "event": event.event_type, "label": f"RTIS {event.event_type} event", "station": event.station, "speed": event.speed or 0, "time": event.event_time.isoformat()}], "home_model": home_model})


@router.get("/rtis/analysis")
def rtis_analysis_page(request: Request, notice: str = "", session: Session = Depends(get_session)):
    uploads = session.exec(
        select(RtisAnalysisUpload).order_by(RtisAnalysisUpload.id.desc()).limit(20)
    ).all()
    home_model, _ = selected_home_model(session)
    return templates.TemplateResponse(
        request=request,
        name="rtis_analysis.html",
        context={
            "active_page": "rtis_analysis",
            "notice": notice,
            "uploads": uploads,
            "home_model": home_model,
        },
    )


@router.post("/rtis/analysis/upload")
def rtis_analysis_upload(
    analysis_date: str = Form(""),
    train_type: str = Form("All"),
    train_no: str = Form(""),
    loco_no: str = Form(""),
    filter_mode: str = Form("route"),
    time_from: str = Form(""),
    time_to: str = Form(""),
    station_from: str = Form(""),
    station_to: str = Form(""),
    gps_file: UploadFile | None = File(None),
    primary_file: UploadFile | None = File(None),
    session: Session = Depends(get_session),
):
    main_file = gps_file if gps_file and gps_file.filename else primary_file
    if main_file is None or not main_file.filename:
        raise HTTPException(400, "Upload a secondary or primary GPS file.")
    filename = (main_file.filename or "gps-data.csv").replace("\\", "/").split("/")[-1]
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx", ".xls"}:
        raise HTTPException(400, "Upload a CSV or Excel GPS data file.")
    try:
        content = main_file.file.read(MAX_BYTES + 1)
        if len(content) > MAX_BYTES:
            raise ValueError(f"File exceeds the {MAX_UPLOAD_MB} MB limit.")
        inferred = _infer_gps_metadata(filename, content)
        primary_content = b""
        primary_filename = ""
        if gps_file and gps_file.filename and primary_file and primary_file.filename:
            primary_filename = (primary_file.filename or "primary-gps.csv").replace("\\", "/").split("/")[-1]
            primary_content = primary_file.file.read(MAX_BYTES + 1)
            if len(primary_content) > MAX_BYTES:
                raise ValueError(f"Primary file exceeds the {MAX_UPLOAD_MB} MB limit.")
        primary_inferred = _infer_gps_metadata(primary_filename, primary_content) if primary_content else {}
        resolved_date = analysis_date.strip() or inferred.get("analysis_date") or primary_inferred.get("analysis_date", "")
        resolved_loco = loco_no.strip()
        if not resolved_loco:
            raise ValueError("Loco no. must be entered manually; the file device ID is not used as loco no.")
        mode = filter_mode.strip().lower() or "route"
        if mode == "time" and (not time_from.strip() or not time_to.strip()):
            raise ValueError("Time mode requires both Time from and Time to.")
        if mode == "station" and (not station_from.strip() or not station_to.strip()):
            raise ValueError("Station mode requires both Station code from and Station code to.")
        if bool(time_from.strip()) != bool(time_to.strip()):
            raise ValueError("Provide both Time from and Time to, or leave both blank.")
        if bool(station_from.strip()) != bool(station_to.strip()):
            raise ValueError("Provide both Station code from and Station code to, or leave both blank.")
        station_from_value = station_from.strip().upper()
        station_to_value = station_to.strip().upper()
        if mode == "station" or station_from_value or station_to_value:
            raw_meta = RtisAnalysisUpload(
                filename=filename, analysis_date="", train_type=train_type.strip(), train_no="",
                loco_no=resolved_loco, time_from="", time_to="", content=content,
                station_from="", station_to="",
            )
            raw_points, _ = _rtis_analysis_points(content, raw_meta)
            raw_codes = {str(point.get("station") or "").strip().upper() for point in raw_points if point.get("station")}
            if primary_content:
                primary_raw_meta = RtisAnalysisUpload(
                    filename=primary_filename, analysis_date="", train_type=train_type.strip(), train_no="",
                    loco_no=resolved_loco, time_from="", time_to="", content=primary_content,
                    station_from="", station_to="",
                )
                primary_raw_points, _ = _rtis_analysis_points(primary_content, primary_raw_meta)
                raw_codes.update(str(point.get("station") or "").strip().upper() for point in primary_raw_points if point.get("station"))
            if not raw_codes:
                raise ValueError("No station code has been found in the raw GPS file.")
            missing = [code for code in (station_from_value, station_to_value) if code not in raw_codes]
            if missing:
                raise ValueError(f"Station code not found in the raw GPS file: {', '.join(missing)}.")
        upload = RtisAnalysisUpload(
            filename=filename, analysis_date=resolved_date, train_type=train_type.strip(),
            train_no=train_no.strip() or primary_inferred.get("train_no") or inferred.get("train_no", ""),
            loco_no=resolved_loco, time_from=time_from.strip(), time_to=time_to.strip(), content=content,
            station_from=station_from_value, station_to=station_to_value,
        )
        session.add(upload)
        session.commit()
        if primary_content:
            session.add(RtisAnalysisPrimaryUpload(analysis_upload_id=upload.id or 0, filename=primary_filename, content=primary_content))
            session.commit()
        notice = f"GPS analysis file saved: {filename}."
    except ValueError as exc:
        session.rollback()
        notice = f"Upload not saved: {exc}"
    finally:
        if gps_file:
            gps_file.file.close()
        if primary_file:
            primary_file.file.close()
    return RedirectResponse("/rtis/analysis?" + urlencode({"notice": notice}), status_code=303)


@router.post("/rtis/analysis/{upload_id}/delete")
def rtis_analysis_delete(upload_id: int, session: Session = Depends(get_session)):
    upload = session.get(RtisAnalysisUpload, upload_id)
    if upload is None:
        return RedirectResponse("/rtis/analysis?" + urlencode({"notice": "Analysis record not found."}), status_code=303)
    primary_uploads = session.exec(
        select(RtisAnalysisPrimaryUpload).where(RtisAnalysisPrimaryUpload.analysis_upload_id == upload_id)
    ).all()
    for primary_upload in primary_uploads:
        session.delete(primary_upload)
    session.delete(upload)
    session.commit()
    return RedirectResponse("/rtis/analysis?" + urlencode({"notice": f"Analysis record deleted: {upload.filename}."}), status_code=303)


def _rtis_analysis_points(content: bytes, upload: RtisAnalysisUpload) -> tuple[list[dict], str]:
    """Read the common SecondaryGPSData CSV columns used by RTIS exports."""
    try:
        if upload.filename.lower().endswith(".xlsx"):
            book = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = book.active
            rows = list(sheet.iter_rows(values_only=True))
            book.close()
            if not rows:
                raise ValueError("The GPS workbook is empty.")
            headers = [str(value or "").strip() for value in rows[0]]
            reader = (dict(zip(headers, row)) for row in rows[1:])
            fieldnames = headers
        else:
            text = content.decode("utf-8-sig", errors="replace")
            csv_reader = csv.DictReader(text.splitlines())
            fieldnames = csv_reader.fieldnames or []
            reader = csv_reader
    except Exception as exc:
        raise ValueError("Cannot read the GPS CSV or Excel file.") from exc
    if not fieldnames:
        raise ValueError("The GPS file has no header row.")
    fields = {str(field).strip().casefold(): field for field in fieldnames if field}
    required = ("logging time", "latitude", "longitude", "speed")
    if any(name not in fields for name in required):
        raise ValueError("GPS file must contain Logging Time, Latitude, Longitude and Speed columns.")
    points: list[dict] = []
    for row in reader:
        try:
            latitude = float(str(row.get(fields["latitude"]) or "").strip())
            longitude = float(str(row.get(fields["longitude"]) or "").strip())
            speed = float(str(row.get(fields["speed"]) or "0").strip() or 0)
        except (TypeError, ValueError):
            continue
        if not (-90 <= latitude <= 90 and -180 <= longitude <= 180):
            continue
        timestamp = str(row.get(fields["logging time"]) or "").strip()
        if upload.analysis_date and not timestamp.startswith(upload.analysis_date):
            continue
        if upload.time_from and timestamp[11:16] < upload.time_from:
            continue
        if upload.time_to and timestamp[11:16] > upload.time_to:
            continue
        station_field = fields.get("last/cur stationcode") or fields.get("station code")
        station = str(row.get(station_field) or "").strip() if station_field else ""
        points.append({"lat": latitude, "lon": longitude, "speed": round(speed, 2), "time": timestamp, "station": station})
    if not points:
        raise ValueError("No GPS points matched the selected filters.")
    return points, ""


def _infer_gps_metadata(filename: str, content: bytes) -> dict[str, str]:
    """Infer date, loco and train number from common GPS export headers/file names."""
    result = {"analysis_date": "", "train_no": "", "loco_no": ""}
    name = Path(filename).stem
    date_match = re.search(r"20\d{2}[-_]\d{2}[-_]\d{2}", name)
    if date_match:
        result["analysis_date"] = date_match.group(0).replace("_", "-")
    try:
        first_line = content.decode("utf-8-sig", errors="replace").splitlines()[0]
        headers = [item.strip().casefold() for item in first_line.split(",")]
        for key, aliases in {
            "analysis_date": ("date", "event date", "train date"),
            "train_no": ("train no", "train number", "trainno"),
        }.items():
            for alias in aliases:
                if alias in headers:
                    index = headers.index(alias)
                    value = content.decode("utf-8-sig", errors="replace").splitlines()[1].split(",")[index].strip()
                    if value:
                        result[key] = value
                    break
    except (IndexError, UnicodeDecodeError):
        pass
    return result


@router.post("/rtis/analysis/run")
def rtis_analysis_run(request: Request, upload_id: int = Form(...), session: Session = Depends(get_session)):
    upload = session.get(RtisAnalysisUpload, upload_id)
    if upload is None:
        raise HTTPException(404, "Analysis upload not found.")
    try:
        points, _ = _rtis_analysis_points(upload.content, upload)
    except ValueError as exc:
        # When no time/station range is supplied, a filename-inferred date may
        # differ from the timestamp format in the GPS export. Read the file
        # without that inferred date filter and let the FSD route bound it.
        if upload.time_from or upload.time_to:
            raise HTTPException(400, str(exc)) from exc
        relaxed_upload = RtisAnalysisUpload(
            filename=upload.filename,
            analysis_date="",
            train_type=upload.train_type,
            train_no=upload.train_no,
            loco_no=upload.loco_no,
            time_from="",
            time_to="",
            content=upload.content,
            station_from=upload.station_from,
            station_to=upload.station_to,
        )
        try:
            points, _ = _rtis_analysis_points(upload.content, relaxed_upload)
        except ValueError as relaxed_exc:
            raise HTTPException(400, str(relaxed_exc)) from relaxed_exc
    primary_upload = session.exec(
        select(RtisAnalysisPrimaryUpload).where(RtisAnalysisPrimaryUpload.analysis_upload_id == upload.id)
    ).first()
    if primary_upload:
        primary_meta = RtisAnalysisUpload(
            filename=primary_upload.filename,
            analysis_date=upload.analysis_date,
            train_type=upload.train_type,
            train_no=upload.train_no,
            loco_no=upload.loco_no,
            time_from=upload.time_from,
            time_to=upload.time_to,
            content=primary_upload.content,
            station_from=upload.station_from,
            station_to=upload.station_to,
        )
        try:
            points.extend(_rtis_analysis_points(primary_upload.content, primary_meta)[0])
        except ValueError:
            pass
    points.sort(key=lambda point: point.get("time", ""))
    home_model, mapping = selected_home_model(session)
    signals = []
    for value in mapping.values():
        for home in value.get("homes", []):
            signals.append({
                "station": value.get("station", ""), "event": value.get("event", ""),
                "label": value.get("label", ""), "line": home.get("line", ""),
                "lat": home.get("latitude"), "lon": home.get("longitude"),
            })
    if not signals:
        raise HTTPException(400, "Upload an FSD signal model before running analysis.")
    if not points:
        raise HTTPException(400, "No GPS points matched the selected date and time period.")

    # Limit the route to the selected FSD station span. With no time or station
    # filters, use the first and last FSD home signal as the default route span.
    signal_indices = []
    for signal in signals:
        if signal.get("lat") is None or signal.get("lon") is None:
            continue
        nearest_index = min(
            range(len(points)),
            key=lambda index: (points[index]["lat"] - signal["lat"]) ** 2
            + (points[index]["lon"] - signal["lon"]) ** 2,
        )
        signal_indices.append((signal, nearest_index))
    station_from = (upload.station_from or "").strip().upper()
    station_to = (upload.station_to or "").strip().upper()
    if station_from and station_to:
        from_indices = [index for signal, index in signal_indices if str(signal.get("station") or "").upper() == station_from]
        to_indices = [index for signal, index in signal_indices if str(signal.get("station") or "").upper() == station_to]
        if not from_indices or not to_indices:
            raise HTTPException(400, "The selected station code was not found in the FSD signal model.")
        start_index, end_index = min(from_indices), max(to_indices)
        if start_index > end_index:
            start_index, end_index = end_index, start_index
        points = points[start_index:end_index + 1]
    elif not upload.time_from and not upload.time_to and signal_indices:
        start_index = min(index for _, index in signal_indices)
        end_index = max(index for _, index in signal_indices)
        points = points[start_index:end_index + 1]
    if not points:
        raise HTTPException(400, "No GPS points matched the selected route.")

    # Match each home signal to the nearest GPS point for its hover details.
    matched_signals = []
    for signal in signals:
        nearest = min(points, key=lambda point: (point["lat"] - signal["lat"]) ** 2 + (point["lon"] - signal["lon"]) ** 2)
        distance_squared = (nearest["lat"] - signal["lat"]) ** 2 + (nearest["lon"] - signal["lon"]) ** 2
        if distance_squared <= 0.001 ** 2:
            signal.update(speed=nearest["speed"], time=nearest["time"])
            matched_signals.append(signal)
    signals = matched_signals
    return templates.TemplateResponse(request=request, name="rtis_analysis_result.html", context={
        "request": request,
        "active_page": "rtis_analysis", "upload": upload, "points": points[::max(1, len(points) // 2000)],
        "signals": signals, "home_model": home_model,
    })


@router.post("/rtis/home-model/upload")
def rtis_home_model_upload(home_file: UploadFile = File(...), division: str = Form("ALL"),
                           return_to: str = Form("rtis"),
                           session: Session = Depends(get_session)):
    if division not in (*DIVISIONS, "ALL"):
        raise HTTPException(400, "Invalid division.")
    try:
        content = home_file.file.read(MAX_BYTES + 1)
        model = save_home_model(session, home_file.filename or '', content, _rtis_ssts_geofences())
        mapped = f"{model.mapped_station_count} station mappings" if model.mapped_station_count else "SSTS station coordinates unavailable; distances will appear after a successful SSTS sync"
        params = {"division": division, "analysis": "passenger",
                  "notice": f"FSD home model saved: {model.signal_count} signals, {mapped}."}
    except ValueError as exc:
        session.rollback()
        params = {"division": division, "analysis": "passenger", "notice": f"FSD home model not saved: {exc}"}
    finally:
        home_file.file.close()
    if return_to == "analysis":
        return RedirectResponse('/rtis/analysis?' + urlencode({"notice": params["notice"]}), status_code=303)
    return RedirectResponse('/rtis?' + urlencode(params), status_code=303)


@router.get("/rtis/home-model.xlsx")
def rtis_home_model_excel(home_station: str = "", session: Session = Depends(get_session)):
    home_station = home_station.strip().upper()
    if len(home_station) > 100:
        raise HTTPException(400, "FSD station search must be 100 characters or fewer.")
    home_model, mapping = selected_home_model(session)
    if home_model is None:
        raise HTTPException(404, "No FSD home signal model has been uploaded.")
    rows = []
    for value in mapping.values():
        for home in value.get("homes", []):
            if home_station and home_station not in value["station"]:
                continue
            rows.append([value["station"], "UP" if value["event"] == "J" else "DOWN", home.get("type", ""),
                         str(value.get("station_latitude", "")), str(value.get("station_longitude", "")),
                         str(home.get("latitude", "")), str(home.get("longitude", "")), home.get("line", ""),
                         str(home.get("distance_m", "")) if home.get("distance_m") is not None else ""])
    rows.sort(key=lambda row: (row[0], row[1], row[5], row[2], row[3], row[4]))
    headers = ("Station", "Direction", "Type", "Station Latitude", "Station Longitude", "Latitude", "Longitude", "Station DIRN", "Linear distance (m)")
    filename = f"RTIS_FSD_Home_Signals_{home_station or 'all-stations'}.xlsx"
    return Response(content=build_rtis_excel(headers, rows),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/rtis/train-model/upload")
def rtis_model_upload(model_file: UploadFile = File(...), division: str = Form("ALL"),
                      session: Session = Depends(get_session)):
    if division not in (*DIVISIONS, "ALL"):
        raise HTTPException(400, "Invalid division.")
    try:
        model = save_train_model(session, model_file.filename or '', model_file.file.read(MAX_BYTES + 1))
        params = {"division": division, "analysis": "passenger", "model_id": model.id,
                  "notice": f"Model saved: {model.eligible_rows} SDAH/KOAA rows, {model.train_count} train numbers. Model filter applied."}
    except ValueError as exc:
        session.rollback()
        params = {"division": division, "analysis": "passenger", "notice": f"Model not saved: {exc}"}
    finally:
        model_file.file.close()
    return RedirectResponse('/rtis?' + urlencode(params), status_code=303)


@router.post("/rtis/upload")
def rtis_upload(division: str = Form(...), files: list[UploadFile] = File(...),
                analysis: str = Form("passenger"),
                session: Session = Depends(get_session)):
    if division not in DIVISIONS or analysis not in ANALYSIS_TYPES:
        raise HTTPException(400, "Choose a valid division.")
    if len(files) > 10:
        raise HTTPException(400, "Upload at most 10 files at a time.")
    messages = []
    for file in files:
        filename = (file.filename or "workbook.xlsx").replace("\\", "/").split("/")[-1]
        try:
            message = import_rtis(session, filename, file.file.read(MAX_BYTES + 1), division)
            if message.startswith(("Saved", "Already uploaded")):
                message = f"Success: {filename} uploaded to {division}. {message}"
        except ValueError as exc:
            session.rollback()
            message = f"Not saved: {exc}"
        finally:
            file.file.close()
        messages.append(f"{filename}: {message}")
    return RedirectResponse("/rtis?" + urlencode({"division": division, "analysis": analysis, "notice": " | ".join(messages)}), status_code=303)


@router.get("/rtis/uploads/{upload_id}/download")
def rtis_download(upload_id: int, session: Session = Depends(get_session)):
    upload = session.get(RtisUpload, upload_id)
    if upload is None:
        raise HTTPException(404, "Upload not found.")
    return Response(content=upload.content,
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="RTIS_{upload.division}_{upload.id}.xlsx"'})


@router.post("/rtis/uploads/{upload_id}/delete")
def rtis_delete_upload(upload_id: int, division: str = Form("ALL"), analysis: str = Form("passenger"),
                       session: Session = Depends(get_session)):
    upload = session.get(RtisUpload, upload_id)
    if upload is None:
        raise HTTPException(404, "Upload not found.")
    deleted_events = session.exec(select(RtisEvent).where(RtisEvent.upload_id == upload_id)).all()
    for event in deleted_events:
        session.delete(event)
    filename = upload.filename
    upload_division = upload.division
    session.delete(upload)
    session.commit()
    return RedirectResponse("/rtis?" + urlencode({
        "division": division if division in (*DIVISIONS, "ALL") else upload_division,
        "analysis": analysis if analysis in ANALYSIS_TYPES else "passenger",
        "notice": f"Undone {filename}: {len(deleted_events):,} events removed.",
    }), status_code=303)


def _output_values(raw: str) -> list[str]:
    source = json.loads(raw)
    return [source.get(header, "") for header in HEADERS]


def output_rows(session: Session):
    """Read saved source fields in a stable order without loading workbook blobs."""
    sources = session.exec(select(RtisEvent.source).order_by(RtisEvent.id)
                           .execution_options(yield_per=1000))
    for raw in sources:
        yield _output_values(raw)


@router.get("/rtis/output")
def rtis_output(request: Request, page: int = 1, session: Session = Depends(get_session)):
    counts = dict(session.exec(select(RtisEvent.division, func.count())
                              .group_by(RtisEvent.division)).all())
    total = sum(counts.values())
    pages = max(1, (total + 99) // 100)
    page = min(max(1, page), pages)
    sources = session.exec(select(RtisEvent.source).order_by(RtisEvent.id)
                           .offset((page - 1) * 100).limit(100)).all()
    return templates.TemplateResponse(request=request, name="rtis_output.html", context={
        "active_page": "rtis", "headers": HEADERS, "divisions": DIVISIONS,
        "counts": counts, "total": total, "page": page, "pages": pages,
        "first_row": (page - 1) * 100 + 1 if total else 0,
        "last_row": min(page * 100, total),
        "rows": [_output_values(raw) for raw in sources],
    })


@router.get("/rtis/output.xlsx")
def rtis_output_excel(session: Session = Depends(get_session)):
    return Response(
        content=build_rtis_excel(HEADERS, output_rows(session)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="rtis_merged_output.xlsx"'},
    )


@router.get("/rtis/output.pdf")
def rtis_output_pdf(session: Session = Depends(get_session)):
    return Response(
        content=build_rtis_pdf(HEADERS, output_rows(session)),
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="rtis_merged_output.pdf"'},
    )
