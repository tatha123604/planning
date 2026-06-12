from __future__ import annotations

from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
from typing import BinaryIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _read_bytes(file_obj) -> bytes:
    if isinstance(file_obj, (bytes, bytearray)):
        return bytes(file_obj)
    if isinstance(file_obj, str):
        with open(file_obj, "rb") as handle:
            return handle.read()
    if hasattr(file_obj, "getvalue"):
        return file_obj.getvalue()
    if hasattr(file_obj, "read"):
        pos = None
        if hasattr(file_obj, "tell"):
            pos = file_obj.tell()
        data = file_obj.read()
        if hasattr(file_obj, "seek") and pos is not None:
            file_obj.seek(pos)
        return data
    raise TypeError("Unsupported file input")


def _as_stream(file_obj) -> BytesIO:
    return BytesIO(_read_bytes(file_obj))


def _normalize_date(value):
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    return value


def infer_report_date(filename: str | None) -> date | None:
    if not filename:
        return None

    stem = Path(filename).stem
    for pattern in (
        r"(?P<day>\d{2})[.\-_](?P<month>\d{2})[.\-_](?P<year>\d{4})",
        r"(?P<year>\d{4})[.\-_](?P<month>\d{2})[.\-_](?P<day>\d{2})",
        r"(?P<day>\d{2})[.\-_](?P<month>\d{2})[.\-_](?P<year>\d{2})",
    ):
        match = re.search(pattern, stem)
        if match:
            try:
                year_text = match.group("year")
                year = int(year_text)
                if len(year_text) == 2:
                    year += 2000
                return date(
                    year,
                    int(match.group("month")),
                    int(match.group("day")),
                )
            except ValueError:
                return None
    return None


def coerce_report_date(value) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%d_%m_%Y", "%d.%m.%y", "%d-%m-%y", "%d_%m_%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def format_report_date(value) -> str | None:
    report_date = coerce_report_date(value)
    if not report_date:
        return None
    return report_date.strftime("%d.%m.%Y")


def format_sheet_title_date(value) -> str | None:
    report_date = coerce_report_date(value)
    if not report_date:
        return None
    return report_date.strftime("%d.%m.%y")


def report_date_iso(value) -> str:
    report_date = coerce_report_date(value)
    if not report_date:
        return ""
    return report_date.isoformat()


def _replace_date_string(value: str, report_date: str) -> str:
    return re.sub(r"\d{2}[.\-_]\d{2}[.\-_]\d{4}", report_date, value)


def _normalize_source_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    df.columns = [str(col).strip().replace("\n", " ") for col in df.columns]
    df = df.rename(
        columns={
            "AllotedDesig.": "Alloted Desig.",
            "Oldest FPOverDue Date": "Oldest FP OverDue Date",
            "CounselOverDue": "Counsel Over Due",
            "Oldest CounselOverDue Date": "Oldest Counsel OverDue Date",
            "GradingOverDue": "Grading OverDue",
            "Oldest GradingOverDue": "Oldest Grading OverDue",
            "TotalOverDue Cases": "Total Over Due Cases",
        }
    )
    needed = [
        "CLI ID",
        "CLI Name",
        "Alloted Desig.",
        "FPOverDue",
        "Oldest FP OverDue Date",
        "Counsel Over Due",
        "Oldest Counsel OverDue Date",
        "Grading OverDue",
        "Oldest Grading OverDue",
    ]
    missing = [col for col in needed if col not in df.columns]
    if missing:
        raise ValueError(f"Missing source columns: {', '.join(missing)}")

    df = df[df["CLI ID"].notna()].copy()
    df["CLI ID"] = df["CLI ID"].astype(str).str.strip()
    df["CLI Name"] = df["CLI Name"].fillna("").astype(str).str.strip()
    df["Alloted Desig."] = df["Alloted Desig."].fillna("").astype(str).str.strip()

    for col in ["FPOverDue", "Counsel Over Due", "Grading OverDue"]:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0).astype(int)

    for col in [
        "Oldest FP OverDue Date",
        "Oldest Counsel OverDue Date",
        "Oldest Grading OverDue",
    ]:
        df[col] = pd.to_datetime(df[col], errors="coerce", dayfirst=True)

    df["Total Over Due Cases"] = (
        df["FPOverDue"] + df["Counsel Over Due"] + df["Grading OverDue"]
    )
    return df


def _load_source_dataframe(source_file) -> pd.DataFrame:
    last_error: Exception | None = None
    for skiprows in (2, 0):
        try:
            df = pd.read_excel(_as_stream(source_file), skiprows=skiprows)
            return _normalize_source_dataframe(df)
        except ValueError as exc:
            last_error = exc
            continue
    if last_error is not None:
        raise last_error
    raise ValueError("Could not read CLI Matrix source workbook.")


def _aggregate_source(df: pd.DataFrame) -> pd.DataFrame:
    grouped = (
        df.groupby(["CLI ID", "CLI Name", "Alloted Desig."], dropna=False, sort=True)
        .agg(
            {
                "FPOverDue": "sum",
                "Oldest FP OverDue Date": "min",
                "Counsel Over Due": "sum",
                "Oldest Counsel OverDue Date": "min",
                "Grading OverDue": "sum",
                "Oldest Grading OverDue": "min",
                "Total Over Due Cases": "sum",
            }
        )
        .reset_index()
    )
    return grouped.sort_values(["CLI ID", "CLI Name", "Alloted Desig."]).reset_index(
        drop=True
    )


def build_summary_df(source_file) -> pd.DataFrame:
    aggregated = _aggregate_source(_load_source_dataframe(source_file))
    summary = aggregated[aggregated["FPOverDue"] > 0].copy()
    summary = summary.rename(columns={"FPOverDue": "FP Over Due"})
    summary = summary[
        [
            "CLI ID",
            "CLI Name",
            "Alloted Desig.",
            "FP Over Due",
            "Oldest FP OverDue Date",
        ]
    ].reset_index(drop=True)
    summary.insert(0, "S.No.", range(1, len(summary) + 1))
    return summary


def build_counselling_summary_df(source_file) -> pd.DataFrame:
    aggregated = _aggregate_source(_load_source_dataframe(source_file))
    summary = aggregated[aggregated["Counsel Over Due"] > 0].copy()
    summary = summary[
        [
            "CLI ID",
            "CLI Name",
            "Alloted Desig.",
            "Counsel Over Due",
            "Oldest Counsel OverDue Date",
        ]
    ].reset_index(drop=True)
    summary.insert(0, "S.No.", range(1, len(summary) + 1))
    return summary


def build_sheet2_df(source_file) -> pd.DataFrame:
    aggregated = _aggregate_source(_load_source_dataframe(source_file)).copy()
    aggregated = aggregated[aggregated["Total Over Due Cases"] > 0].copy()
    aggregated = aggregated.rename(columns={"FPOverDue": "FP Over Due"})
    aggregated = aggregated[
        [
            "CLI ID",
            "CLI Name",
            "Alloted Desig.",
            "FP Over Due",
            "Oldest FP OverDue Date",
            "Counsel Over Due",
            "Oldest Counsel OverDue Date",
            "Grading OverDue",
            "Oldest Grading OverDue",
            "Total Over Due Cases",
        ]
    ].reset_index(drop=True)
    aggregated.insert(0, "S.No.", range(1, len(aggregated) + 1))
    return aggregated


def _copy_cell_style(source_cell, target_cell) -> None:
    target_cell._style = copy(source_cell._style)
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


def _clear_merges(ws, from_row: int) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= from_row:
            ws.unmerge_cells(str(merged_range))


def _clear_range(ws, start_row: int, end_row: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(1, end_col + 1):
            ws.cell(row=row, column=col).value = None


def _apply_row_style(ws, source_row: int, target_row: int, end_col: int) -> None:
    for col in range(1, end_col + 1):
        _copy_cell_style(ws.cell(source_row, col), ws.cell(target_row, col))
    if ws.row_dimensions[source_row].height is not None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _write_dataframe(ws, df: pd.DataFrame, start_row: int) -> None:
    for row_index in range(len(df)):
        for col_index, col_name in enumerate(df.columns, start=1):
            value = df.iloc[row_index, col_index - 1]
            ws.cell(row=start_row + row_index, column=col_index).value = _normalize_date(value)


def _merge_same_cli(ws, start_row: int, count: int, cli_col: int, name_col: int) -> None:
    if count <= 1:
        return
    block_start = start_row
    current_cli = ws.cell(start_row, cli_col).value
    for row in range(start_row + 1, start_row + count):
        value = ws.cell(row, cli_col).value
        if value != current_cli:
            if current_cli and row - 1 > block_start:
                ws.merge_cells(
                    start_row=block_start,
                    end_row=row - 1,
                    start_column=cli_col,
                    end_column=cli_col,
                )
                ws.merge_cells(
                    start_row=block_start,
                    end_row=row - 1,
                    start_column=name_col,
                    end_column=name_col,
                )
            block_start = row
            current_cli = value
    if current_cli and start_row + count - 1 > block_start:
        ws.merge_cells(
            start_row=block_start,
            end_row=start_row + count - 1,
            start_column=cli_col,
            end_column=cli_col,
        )
        ws.merge_cells(
            start_row=block_start,
            end_row=start_row + count - 1,
            start_column=name_col,
            end_column=name_col,
        )


def _find_total_row(ws, start_row: int) -> int:
    for row in range(start_row, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str) and "total" in value.lower():
            return row
    return start_row + 19


def _trim_trailing_empty_rows(ws, start_row: int, end_col: int) -> None:
    last_used_row = start_row - 1
    for row in range(start_row, ws.max_row + 1):
        if any(ws.cell(row, col).value not in (None, "") for col in range(1, end_col + 1)):
            last_used_row = row
    if last_used_row < ws.max_row:
        ws.delete_rows(last_used_row + 1, ws.max_row - last_used_row)


def _update_sheet_headers(ws, report_date: str | None) -> None:
    if not report_date:
        return

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "CLI MATRIX" in cell.value.upper():
                cell.value = _replace_date_string(cell.value, report_date)

    for header_name in ("oddHeader", "evenHeader", "firstHeader"):
        header = getattr(ws, header_name, None)
        if header is None:
            continue
        for part_name in ("left", "center", "right"):
            part = getattr(header, part_name, None)
            if part is not None and getattr(part, "text", ""):
                part.text = _replace_date_string(part.text, report_date)


def _write_summary_sheet(
    ws,
    df: pd.DataFrame,
    end_col: int,
    total_label: str,
    total_col_index: int,
) -> None:
    data_start = 3
    total_template_row = _find_total_row(ws, data_start)
    _clear_merges(ws, data_start)
    _clear_range(ws, data_start, max(ws.max_row, data_start + len(df) + 2), end_col)

    for offset in range(len(df)):
        _apply_row_style(ws, 3, data_start + offset, end_col)

    _write_dataframe(ws, df, data_start)
    _merge_same_cli(ws, data_start, len(df), 2, 3)

    total_row = data_start + len(df)
    _apply_row_style(ws, total_template_row, total_row, end_col)
    ws.cell(total_row, 1).value = total_label
    total_col_letter = get_column_letter(total_col_index)
    ws.cell(total_row, total_col_index).value = f"=SUM({total_col_letter}{data_start}:{total_col_letter}{total_row - 1})"
    ws.cell(total_row, 2).value = None
    ws.cell(total_row, 3).value = None
    ws.cell(total_row, 4).value = None
    for col in range(total_col_index + 1, end_col + 1):
        ws.cell(total_row, col).value = None
    ws.merge_cells(start_row=total_row, end_row=total_row, start_column=1, end_column=4)
    ws.merge_cells(
        start_row=total_row,
        end_row=total_row,
        start_column=total_col_index,
        end_column=end_col,
    )
    _trim_trailing_empty_rows(ws, data_start, end_col)


def _write_sheet1(ws, df: pd.DataFrame) -> None:
    _write_summary_sheet(ws, df, 6, "TOTAL FP DUE", 5)


def _write_counselling_sheet(ws, df: pd.DataFrame) -> None:
    _write_summary_sheet(ws, df, 6, "TOTAL COUNSELLING DUE", 5)


def _write_sheet2(ws, df: pd.DataFrame) -> None:
    data_start = 3
    end_col = 11
    _clear_merges(ws, data_start)
    _clear_range(ws, data_start, max(ws.max_row, data_start + len(df) + 2), end_col)

    for offset in range(len(df)):
        _apply_row_style(ws, 3, data_start + offset, end_col)

    _write_dataframe(ws, df, data_start)
    _merge_same_cli(ws, data_start, len(df), 2, 3)
    _trim_trailing_empty_rows(ws, data_start, end_col)


def _get_sheet_by_name_or_pattern(workbook, preferred_names, pattern=None):
    for name in preferred_names:
        if name in workbook.sheetnames:
            return workbook[name]
    if pattern:
        for ws in workbook.worksheets:
            if re.fullmatch(pattern, ws.title):
                return ws
    raise ValueError(
        "Template workbook is missing one of the required report sheets."
    )


def build_output_workbook(
    source_file,
    template_file,
    source_filename: str | None = None,
    report_date_value=None,
) -> BytesIO:
    summary_df = build_summary_df(source_file)
    counselling_df = build_counselling_summary_df(source_file)
    sheet2_df = build_sheet2_df(source_file)
    report_date = format_report_date(report_date_value) or format_report_date(
        infer_report_date(source_filename)
    )
    sheet_title_date = format_sheet_title_date(infer_report_date(source_filename)) or format_sheet_title_date(
        report_date_value
    )

    workbook = load_workbook(_as_stream(template_file))
    sheet1 = _get_sheet_by_name_or_pattern(
        workbook,
        ["Summary position of FP OVERDUE"],
        r"\d{2}[.\-_]\d{2}[.\-_]\d{2,4}",
    )
    sheet2 = _get_sheet_by_name_or_pattern(
        workbook,
        ["COUNSELLING DUE SUMMARY"],
        r"COUNSELLING DUE SUMMARY",
    )
    sheet3 = _get_sheet_by_name_or_pattern(
        workbook,
        ["13.04.26"],
        r"\d{2}[.\-_]\d{2}[.\-_]\d{2}",
    )

    _write_sheet1(sheet1, summary_df)
    _write_counselling_sheet(sheet2, counselling_df)
    _write_sheet2(sheet3, sheet2_df)
    _update_sheet_headers(sheet1, report_date)
    _update_sheet_headers(sheet2, report_date)
    _update_sheet_headers(sheet3, report_date)
    if sheet_title_date:
        sheet3.title = sheet_title_date
    workbook.calculation.fullCalcOnLoad = True

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
